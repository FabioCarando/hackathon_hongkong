"""Generate the fake finance workspace for the Trace demo (Lantau Peak Family Office Ltd).

Run:
    uv run --with openpyxl --with fpdf2 --with pymupdf python scripts/generate_workspace.py

Deterministic (fixed seed, fixed timestamps). Wipes and rebuilds workspace/, including its own git
history with backdated commits. That history lives in workspace/.trace/git (not workspace/.git) so the
outer repo can track the workspace files and its history as plain files. Use it with:
    git --git-dir=workspace/.trace/git log      (core.worktree points back at workspace/)
Inbox documents, the 2027 lease and the two newest emails stay untracked (they are the demo).
ground_truth/ is git-ignored inside the workspace.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import random
import re
import shutil
import subprocess
import zipfile
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from email.policy import SMTPUTF8
from pathlib import Path

import pymupdf as fitz
from fontTools.ttLib import TTCollection
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from PIL import Image, ImageFilter

ROOT = Path(__file__).resolve().parents[1]
WS = ROOT / "workspace"
GIT_DIR = WS / ".trace" / "git"
SEED = 20261005
rng = random.Random(SEED)
TODAY = date(2026, 10, 5)
CLOSE = date(2026, 9, 30)
D = date

# ----------------------------------------------------------------------------- master data
CO = dict(
    name_en="Lantau Peak Family Office Ltd", name_zh="大嶼峰家族辦公室有限公司", br="74120583",
    addr_en="Suite 3802, 38/F, Harbour Crest Tower, 1 Harbour View Street, Central, Hong Kong",
    addr_zh="香港中環海景街1號海冠大廈38樓3802室",
    bank="Victoria Harbour Bank Limited", bank_acct="088-517-22046-8", domain="lantaupeak.com.hk",
    family="Cheung family", trust="Cheung Family Trust", trust_zh="張氏家族信託",
)
PEOPLE = {"victoria": ("Victoria Cheung", "victoria.cheung@lantaupeak.com.hk", "Principal (family member)"),
          "raymond": ("Raymond Ho", "raymond.ho@lantaupeak.com.hk", "Chief Financial Officer"),
          "grace": ("Grace Lam", "grace.lam@lantaupeak.com.hk", "Fund Accountant"),
          "jason": ("Jason Yip", "jason.yip@lantaupeak.com.hk", "Accounts Clerk")}
APPROVAL_LIMIT = 500_000

SUP = {
    "S01": dict(name_en="Pearl River Growth Fund II LP", name_zh="珠江成长基金二期", currency="RMB", terms=14,
                domain="prgfund.com", email="investor.services@prgfund.com", account="150",
                bank="Pearl River Commercial Bank, Shenzhen Futian Branch (珠江商贸银行深圳福田支行)",
                bank_acct="6214 8320 0019 2049", reg="GP: Pearl River Capital Management (Shenzhen) Co., Ltd.",
                tel="+86 755 2861 4402", addr_zh="深圳市福田区益田路6001号太平金融大厦18楼",
                addr_en="18/F, Taiping Finance Tower, 6001 Yitian Road, Futian, Shenzhen",
                contract="docs/contracts/S01_subscription_agreement_2024.pdf", contract_no="PRG2-SUB-0042",
                commitment=50_000_000.0, called_before=14_000_000.0, kind="RMB private equity fund (Greater Bay Area growth)"),
    "S02": dict(name_en="Harbourview Capital Partners III LP", name_zh="", currency="USD", terms=10,
                domain="harbourviewcap.com", email="fundadmin@harbourviewcap.com", account="150",
                bank="Atlantic Mercantile Bank, New York (SWIFT ATMBUS33)", bank_acct="8841-2207-5530",
                reg="Cayman Islands exempted limited partnership, reg. no. MC-338172", tel="+852 3958 2100",
                addr_en="Level 21, 8 Finance Street West, Central, Hong Kong (Manager: Harbourview Capital Management Ltd)",
                contract="docs/contracts/S02_side_letter_2024.pdf", contract_no="HCP3-SL-017",
                commitment=10_000_000.0, called_before=3_800_000.0, kind="USD buyout fund (Asia mid-market)"),
    "S03": dict(name_en="Peak Estates Property Management Ltd", name_zh="峰譽物業管理有限公司", currency="HKD", terms=30,
                domain="peakestates.com.hk", email="accounts@peakestates.com.hk", account="471",
                bank="Victoria Harbour Bank Limited", bank_acct="088-221-55190-3", reg="BR No. 50993312",
                tel="+852 2418 6630", addr_en="Unit 1608, Wanchai Commercial Centre, 194 Johnston Road, Wan Chai, Hong Kong",
                contract="docs/contracts/S03_property_management_agreement_2026.pdf"),
    "S04": dict(name_en="Halcyon Re Asia Ltd", name_zh="", currency="HKD", terms=0, domain="halcyonre.com.hk",
                email="facilities@halcyonre.com.hk", account="260", bank="Lion Rock Bank Limited",
                bank_acct="512-339-80417-0", reg="BR No. 21874460", tel="+852 2795 3100",
                addr_en="27/F, Quarry Bay Exchange, 979 King's Road, Quarry Bay, Hong Kong",
                contract="docs/contracts/S04_lease_2024_2026.pdf"),
    "S05": dict(name_en="Meridian Fine Art Insurance Ltd", name_zh="", currency="USD", terms=15,
                domain="meridianfineart.com", email="billing@meridianfineart.com", account="433",
                bank="Sierra Mesa Bank (SWIFT SMBKUS6L)", bank_acct="4410-0928-1173", reg="Bermuda reg. no. 58821",
                tel="+1 441 555 0142", addr_en="Clarendon House, 2 Church Street, Hamilton HM 11, Bermuda",
                contract="docs/contracts/S05_art_insurance_policy_2025.pdf"),
}
FAKE_DOMAIN = "prg-fund.co"
NEW_BANK = ("Nanhai Union Bank, Shenzhen Bao'an Branch (南海联合银行深圳宝安支行)", "6230 5821 4407 7731")
FEE_RATE, FEE_RATE_WRONG = 1.50, 2.00

PM_RATES = {  # Peak Estates property management agreement, schedule 1
    "PM-RB12": ("Property management - Flat 12A, Seaview Court, Repulse Bay", "month", 8800),
    "PM-CR21": ("Property management - Flat 21B, Pinecrest Tower, Mid-Levels", "month", 7600),
    "KEY-HLD": ("Key holding & contractor access (per flat)", "month", 450),
    "LS-INSP": ("Tenancy inspection with photo report", "visit", 1200),
    "CLN-DP": ("Deep cleaning, vacant areas & balcony", "visit", 3500),
}
PROPS = {"12A": "Flat 12A, 12/F, Seaview Court, 38 Repulse Bay Road, Repulse Bay, Hong Kong",
         "21B": "Flat 21B, 21/F, Pinecrest Tower, 9 Conduit Road, Mid-Levels, Hong Kong"}

FX = {"2025-09": (7.79, 1.093), "2025-10": (7.78, 1.090), "2025-11": (7.79, 1.074), "2025-12": (7.78, 1.071),
      "2026-01": (7.79, 1.072), "2026-02": (7.81, 1.077), "2026-03": (7.82, 1.079), "2026-04": (7.80, 1.076),
      "2026-05": (7.79, 1.073), "2026-06": (7.81, 1.078), "2026-07": (7.82, 1.083), "2026-08": (7.81, 1.082),
      "2026-09": (7.80, 1.080), "2026-10": (7.80, 1.080)}

ACCOUNTS = [  # code, name, type, group
    ("090", "Victoria Harbour Bank - HKD Current", "Bank", "Balance sheet"),
    ("150", "Investments - Private Funds (at cost)", "Non-current Asset", "Balance sheet"),
    ("151", "Investments - Listed Portfolio (custodian)", "Non-current Asset", "Balance sheet"),
    ("160", "Investment Property", "Non-current Asset", "Balance sheet"),
    ("620", "Prepayments", "Current Asset", "Balance sheet"),
    ("800", "Accounts Payable", "Current Liability", "Balance sheet"),
    ("820", "Accrued Liabilities", "Current Liability", "Balance sheet"),
    ("825", "MPF Payable", "Current Liability", "Balance sheet"),
    ("830", "Tenant Deposits Held", "Current Liability", "Balance sheet"),
    ("960", "Family Capital & Retained Earnings", "Equity", "Balance sheet"),
    ("970", "Share Capital", "Equity", "Balance sheet"),
    ("260", "Rental Income", "Revenue", "Income"),
    ("270", "Interest Income", "Revenue", "Income"),
    ("275", "Dividend & Portfolio Income", "Revenue", "Income"),
    ("477", "Salaries - Office & Household Staff", "Expense", "Staff"),
    ("478", "MPF Contributions", "Expense", "Staff"),
    ("471", "Property Management Fees", "Expense", "Property"),
    ("473", "Property Repairs & Maintenance", "Expense", "Property"),
    ("475", "Building Management & Government Rates", "Expense", "Property"),
    ("481", "School & Tuition Fees", "Expense", "Education"),
    ("493", "Travel & Aviation Charter", "Expense", "Travel & aviation"),
    ("433", "Insurance", "Expense", "Insurance"),
    ("412", "Legal, Tax & Audit Fees", "Expense", "Professional fees"),
    ("455", "Fund Management Fees", "Expense", "Fund fees"),
    ("404", "Bank Fees", "Expense", "Admin"),
    ("429", "General Office Expenses", "Expense", "Admin"),
    ("485", "Software Subscriptions", "Expense", "Admin"),
    ("489", "Telephone & Internet", "Expense", "Admin"),
    ("497", "Foreign Exchange Gain/Loss", "Expense", "Admin"),
]
ACC = {a[0]: a for a in ACCOUNTS}
PL_CODES = [a[0] for a in ACCOUNTS if a[3] != "Balance sheet"]
GROUPS = ["Income", "Staff", "Property", "Education", "Travel & aviation", "Insurance", "Professional fees",
          "Fund fees", "Admin"]
BUDGET = {"260": 157_000, "270": 85_000, "275": 3_800_000, "477": 420_000, "478": 15_000, "471": 18_000,
          "473": 3_000, "475": 11_600, "481": 0, "493": 125_000, "433": 19_000, "412": 28_000, "455": 0,
          "404": 900, "429": 9_000, "485": 3_200, "489": 2_400, "497": 0}
BUDGET_ONE_OFF = {("481", 1): 286_000, ("481", 8): 312_000, ("433", 1): 48_000, ("412", 3): 120_000,
                  ("455", 1): 292_500, ("455", 4): 292_500, ("455", 7): 292_500, ("455", 10): 292_500}


def budget(code, m):
    return BUDGET.get(code, 0) + BUDGET_ONE_OFF.get((code, m), 0)


# ----------------------------------------------------------------------------- helpers
def bday(d):
    while d.weekday() >= 5:
        d += timedelta(1)
    return d


def ym(d):
    return d.strftime("%Y-%m")


def fx(cur, d):
    return 1.0 if cur == "HKD" else FX[ym(d)][0 if cur == "USD" else 1]


def r2(x):
    return round(x + 1e-9, 2)


def money(x):
    return f"{x:,.2f}"


def months(a, b):
    y, m = a
    while (y, m) <= b:
        yield y, m
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)


def month_end(y, m):
    return D(y + (m == 12), m % 12 + 1, 1) - timedelta(1)


def last_bday(y, m):
    d = month_end(y, m)
    while d.weekday() >= 5:
        d -= timedelta(1)
    return d


def split2(s, n):
    """Greedy-wrap an address at ', ' boundaries into lines of at most ~n chars."""
    lines = [""]
    for part in s.split(", "):
        if lines[-1] and len(lines[-1]) + len(part) + 2 > n:
            lines.append("")
        lines[-1] += (", " if lines[-1] else "") + part
    return [x + ("," if i < len(lines) - 1 else "") for i, x in enumerate(lines)]


def zh_date(d):
    return f"{d.year}年{d.month}月{d.day}日"


def en_date(d):
    return f"{d.day} {d:%b %Y}"


# ----------------------------------------------------------------------------- payable documents
INVOICES: list[dict] = []


def mk_inv(sid, no, d, lines, file, folder, kind="invoice", bank=None, email=None, account=None, terms=None, **extra):
    s = SUP[sid]
    sub = r2(sum(r2(q * u) for _, _, q, u in lines))
    inv = dict(sid=sid, no=no, date=d, due=d + timedelta(s["terms"] if terms is None else terms), currency=s["currency"],
               lines=lines, subtotal=sub, tax=0.0, total=sub, account=account or s["account"], kind=kind,
               bank=bank or s["bank_acct"], bank_name=s["bank"] if bank is None else NEW_BANK[0],
               email=email or s["email"], file=file, folder=folder, **extra)
    inv["pay_date"] = bday(inv["due"])
    INVOICES.append(inv)
    return inv


def folder_of(d):
    return f"docs/invoices/{ym(d)}" if d >= D(2026, 7, 1) else None


# S01: Pearl River Growth Fund II drawdowns (every two months), RMB
S01_CALLS = [(D(2025, 9, 8), 12, 2_500_000), (D(2025, 11, 10), 13, 3_000_000), (D(2026, 1, 8), 14, 2_000_000),
             (D(2026, 3, 9), 15, 3_500_000), (D(2026, 5, 8), 16, 2_500_000), (D(2026, 7, 8), 17, 3_000_000)]
called = SUP["S01"]["called_before"]
for d, n, amt in S01_CALLS:
    no = f"PRG2-DN-{n:03d}"
    mk_inv("S01", no, d, [(f"DD-{n:03d}", f"第{n}次缴款 Drawdown No. {n}", 1, float(amt))], f"Pearl River call {no}.pdf",
           folder_of(d), kind="capital_call", called_before=called)
    called += amt

# S02: Harbourview Capital Partners III: quarterly management fee notices + capital calls, USD
S02_FEES = [(D(2025, 10, 1), "2025Q4"), (D(2026, 1, 2), "2026Q1"), (D(2026, 4, 1), "2026Q2"), (D(2026, 7, 1), "2026Q3")]
QTR = {"Q1": ("1 Jan", "31 Mar"), "Q2": ("1 Apr", "30 Jun"), "Q3": ("1 Jul", "30 Sep"), "Q4": ("1 Oct", "31 Dec")}


def fee_line(period, rate):
    y, q = period[:4], period[4:]
    amt = r2(SUP["S02"]["commitment"] * rate / 100 / 4)
    return [("MGMT-FEE", f"Management fee {QTR[q][0]} - {QTR[q][1]} {y}: {rate:.2f}% p.a. x commitment USD 10,000,000 x 1/4", 1, amt)]


for d, period in S02_FEES:
    no = f"HCP3-MF-{period}"
    mk_inv("S02", no, d, fee_line(period, FEE_RATE), f"Harbourview fee {no}.pdf", folder_of(d), kind="fee_notice",
           account="455", terms=30, rate=FEE_RATE, period=period)
S02_CALLS = [(D(2026, 2, 16), 1, 400_000), (D(2026, 5, 15), 2, 500_000), (D(2026, 8, 14), 3, 500_000)]
called = SUP["S02"]["called_before"]
for d, n, amt in S02_CALLS:
    no = f"HCP3-CN-2026-{n:02d}"
    mk_inv("S02", no, d, [(f"CALL-{n:02d}", f"Capital call {n}/2026: {amt / 1e5:.0f}% of commitment (investments and fund expenses)", 1, float(amt))],
           f"Harbourview call {no}.pdf", folder_of(d), kind="capital_call", called_before=called)
    called += amt


# S03: Peak Estates monthly property management invoices, HKD
def pm_lines(insp, clean):
    q = {"PM-RB12": 1, "PM-CR21": 1, "KEY-HLD": 2, "LS-INSP": insp, "CLN-DP": clean}
    return [(k, PM_RATES[k][0], q[k], float(PM_RATES[k][2])) for k in PM_RATES if q[k]]


S03_NO = {(2026, 7): 2204, (2026, 8): 2247, (2026, 9): 2291}
for k, (y, m) in enumerate(months((2025, 11), (2026, 9))):
    d = bday(D(y, m, 1 if (y, m) == (2026, 9) else 3))
    no = f"INV-PM-{S03_NO.get((y, m), 1984 + 22 * k)}"
    insp, clean = ((2, 1) if (y, m) == (2026, 9) else (rng.randint(0, 2), rng.randint(0, 1)))
    mk_inv("S03", no, d, pm_lines(insp, clean), f"Peak Estates {no}.pdf", folder_of(d))

# S05: Meridian Fine Art Insurance monthly premium instalments, USD
for k, (y, m) in enumerate(months((2025, 11), (2026, 9))):
    d = bday(D(y, m, 1))
    no = f"MAI-{10011 + 11 * k}"
    mk_inv("S05", no, d, [("ART-PREM", f"Fine art collection policy MAI-FA-2025-0716, premium instalment ({d:%b %Y})", 1, 2450.0)],
           f"Meridian {no}.pdf", folder_of(d))

HIST = [i for i in INVOICES if i["folder"]]  # Jul-Sep 2026, in the register
by_no = {i["no"]: i for i in INVOICES}
INV2291 = by_no["INV-PM-2291"]
INV2291["entered"] = D(2026, 9, 3)

INBOX = [
    mk_inv("S01", "PRG2-DN-018", D(2026, 9, 30), [("DD-018", "第18次缴款 Drawdown No. 18", 1, 3_500_000.0)],
           "scan_1002.pdf", "docs/invoices/inbox", kind="capital_call", bank=NEW_BANK[1], email=f"accounts@{FAKE_DOMAIN}",
           called_before=SUP["S01"]["called_before"] + sum(a for *_, a in S01_CALLS)),
    mk_inv("S02", "HCP3-CN-2026-04", D(2026, 9, 30),
           [("CALL-04", "Capital call 4/2026: 6% of commitment (investments and fund expenses)", 1, 600_000.0)],
           "HCP3_Capital_Call_Notice_04.pdf", "docs/invoices/inbox", kind="capital_call",
           called_before=SUP["S02"]["called_before"] + sum(a for *_, a in S02_CALLS)),
    mk_inv("S02", "HCP3-MF-2026Q4", D(2026, 10, 1), fee_line("2026Q4", FEE_RATE_WRONG), "HCP3_Q4_2026_Management_Fee.pdf",
           "docs/invoices/inbox", kind="fee_notice", account="455", terms=30, rate=FEE_RATE_WRONG, period="2026Q4"),
    mk_inv("S03", "INV-PM-2291-R", D(2026, 9, 25), INV2291["lines"], "Invoice (3).pdf", "docs/invoices/inbox"),
    mk_inv("S03", "INV-PM-2318", D(2026, 9, 30), pm_lines(1, 0), "scan_0930_2.pdf", "docs/invoices/inbox"),
    mk_inv("S05", "MAI-10044", D(2026, 10, 1),
           [("ART-PREM", "Fine art collection policy MAI-FA-2025-0716, premium instalment (Oct 2026)", 1, 2450.0)],
           "Invoice_MAI-10044.pdf", "docs/invoices/inbox"),
]
for i in INBOX:
    INVOICES.remove(i)  # not entered anywhere yet
SCANNED = {"scan_1002.pdf", "scan_0930_2.pdf"}

# ----------------------------------------------------------------------------- general ledger
JOURNALS: list[dict] = []


def jnl(d, source, ref, contact, desc, lines, bank_desc=None, cat=None):
    JOURNALS.append(dict(date=d, source=source, ref=ref, contact=contact, desc=desc, bank_desc=bank_desc, cat=cat,
                         lines=[(a, r2(dr), r2(cr)) for a, dr, cr in lines if dr or cr]))


CF_CAT = {"150": "Fund capital calls", "455": "Fund management fees", "471": "Property", "433": "Insurance"}
PAYMENTS = []
Y0 = D(2026, 1, 1)
open_ap = 0.0
for inv in INVOICES:
    s = SUP[inv["sid"]]
    booked = r2(inv["total"] * fx(inv["currency"], inv["date"]))
    inv["amount_hkd"] = booked
    if inv["date"] >= Y0:
        jnl(inv["date"], "Payable Invoice", inv["no"], s["name_en"], f"{s['name_en']} {inv['no']}",
            [(inv["account"], booked, 0), ("800", 0, booked)])
    if inv["pay_date"] > CLOSE:
        if inv["date"] < Y0:
            open_ap += booked
        continue
    paid = r2(inv["total"] * fx(inv["currency"], inv["pay_date"]))
    PAYMENTS.append(dict(payment_date=inv["pay_date"], supplier_id=inv["sid"], supplier=s["name_en"],
                         invoice_no=inv["no"], amount=inv["total"], currency=inv["currency"], amount_hkd=paid,
                         bank_name=s["bank"].split(" (")[0], bank_account_paid=s["bank_acct"],
                         payment_ref=f"PAY-{inv['pay_date']:%y%m%d}-{inv['sid']}"))
    if inv["pay_date"] < Y0:
        continue
    if inv["date"] < Y0:
        open_ap += booked
    diff = r2(paid - booked)
    method = "TT OUT" if inv["currency"] != "HKD" else "FPS OUT"
    jnl(inv["pay_date"], "Payable Payment", inv["no"], s["name_en"], f"Payment {inv['no']}",
        [("800", booked, 0), ("497", max(diff, 0), max(-diff, 0)), ("090", 0, paid)],
        f"{method} {s['name_en'].upper().rstrip('.')} {inv['no']}", CF_CAT[inv["account"]])
    if inv["currency"] != "HKD":  # TT charge: own bank line, same journal
        JOURNALS[-1]["lines"] += [("404", 150.0, 0.0), ("090", 0.0, 150.0)]
        JOURNALS[-1]["bank_desc2"] = "TT CHARGE"

OPEN_CASH = 18_500_000.00
DEPOSITS = 285_000 + 186_000
ob = [("090", OPEN_CASH, 0), ("150", 152_400_000, 0), ("151", 2_350_000_000, 0), ("160", 186_000_000, 0),
      ("620", 96_000, 0), ("800", 0, open_ap), ("830", 0, DEPOSITS), ("970", 0, 10_000)]
plug = r2(sum(l[1] - l[2] for l in ob))
assert plug > 0
jnl(Y0, "Manual Journal", "OB-2026", "", "Opening balances b/f 31 Dec 2025", ob + [("960", 0, plug)])

# income: rent, portfolio income sweep, interest
for y, m in months((2026, 1), (2026, 9)):
    d = bday(D(y, m, 1))
    jnl(d, "Receive Money", f"RNT-{m:02d}-12A", SUP["S04"]["name_en"], f"Rent Flat 12A Seaview Court ({d:%b %Y})",
        [("090", 95000, 0), ("260", 0, 95000)], f"CR TRF HALCYON RE ASIA LTD RENT {d:%b%y}".upper(), "Rental income")
    d2 = bday(D(y, m, 2))
    jnl(d2, "Receive Money", f"RNT-{m:02d}-21B", "Brightwater Consulting Ltd", f"Rent Flat 21B Pinecrest Tower ({d:%b %Y})",
        [("090", 62000, 0), ("260", 0, 62000)], f"CR TRF BRIGHTWATER CONSULTING RENT {d:%b%y}".upper(), "Rental income")
    sweep = float(rng.randint(3400, 4300) * 1000)
    jnl(bday(D(y, m, 25)), "Receive Money", f"SWP-{m:02d}", "Lion Rock Private Bank (custodian)",
        "Portfolio income sweep (dividends and coupons)", [("090", sweep, 0), ("275", 0, sweep)],
        "CR TRF LION ROCK PB INCOME SWEEP", "Portfolio income")

# recurring spend money (vendor, account, amount fn, day, bank desc, cash flow category)
SPEND = [
    ("Seaview Court & Pinecrest Tower management offices", "475", lambda m: 11600 if m < 6 else 12800, 2,
     "AUTOPAY BLDG MGMT SEAVIEW/PINECREST", "Property"),
    ("Harbourtel Broadband & Mobile", "489", lambda m: 2380, 5, "DD HARBOURTEL", "Office & admin"),
    ("Island Wide Handyman Services", "473", lambda m: rng.randint(10, 28) * 100, 18, "FPS OUT ISLAND WIDE HANDYMAN", "Property"),
    ("SkyJet Charter Asia (travel desk)", "493", lambda m: rng.randint(80, 170) * 1000, 20, "FPS OUT SKYJET CHARTER ASIA", "Family travel"),
    ("Central Office Supplies Co", "429", lambda m: rng.randint(60, 110) * 100, 22, "EPS CENTRAL OFFICE SUPPLIES", "Office & admin"),
    ("Wealth platform & software licences (card)", "485", lambda m: 3150, 3, "CARD SOFTWARE LICENCES", "Office & admin"),
    ("Kwan & Partners Tax Advisory", "412", lambda m: 28000, 25, "FPS OUT KWAN & PARTNERS", "Professional fees"),
]
ONE_OFF = [
    (D(2026, 1, 12), "Peak Assurance Ltd", "433", 48000, "Household & contents policy 2026 - family residence",
     "CHQ PEAK ASSURANCE", "Insurance"),
    (D(2026, 1, 15), "Island International School", "481", 286000, "Spring term 2026 tuition - 2 children",
     "FPS OUT ISLAND INTL SCHOOL", "Education"),
    (D(2026, 3, 20), "Lam & Partners CPA", "412", 120000, "2025 audit and tax compliance", "FPS OUT LAM & PARTNERS CPA",
     "Professional fees"),
    (D(2026, 7, 14), "Swiftfix Building Services", "473", 38500,
     "Emergency water-pipe repair after typhoon - Flat 12A Seaview Court", "FPS OUT SWIFTFIX BLDG SERVICES", "Property"),
    (D(2026, 8, 17), "Island International School", "481", 318000, "Autumn term 2026 tuition - 2 children (fees +2%)",
     "FPS OUT ISLAND INTL SCHOOL", "Education"),
    (D(2026, 9, 16), "Arctic Breeze Engineering Co", "473", 9800, "Air-con replacement - Flat 21B Pinecrest Tower",
     "FPS OUT ARCTIC BREEZE ENG", "Property"),
]
for y, m in months((2026, 1), (2026, 9)):
    for vendor, acct, fn, day, bd, cat in SPEND:
        amt = float(fn(m))
        desc = vendor + (" (revised fee from Jun 2026)" if acct == "475" and m >= 6 else "")
        jnl(bday(D(y, m, day)), "Spend Money", f"SM-{m:02d}-{acct}", vendor, desc, [(acct, amt, 0), ("090", 0, amt)], bd, cat)
    pd = last_bday(y, m)
    jnl(pd, "Spend Money", f"PAY-{m:02d}", "Payroll", f"Payroll {pd:%b %Y} (4 office + 6 household staff)",
        [("477", 420000, 0), ("478", 15000, 0), ("090", 0, 405000), ("825", 0, 30000)], f"PAYROLL AUTOPAY {pd:%b%y}".upper(),
        "Staff & MPF")
    jnl(pd, "Spend Money", f"MPF-{m:02d}", "Orchid Trust MPF Scheme", "MPF contributions (ER + EE)",
        [("825", 30000, 0), ("090", 0, 30000)], "MPF ORCHID TRUST", "Staff & MPF")
    interest = float(rng.randint(70, 100) * 1000 + rng.randint(0, 999))
    jnl(pd, "Receive Money", f"INT-{m:02d}", CO["bank"], "Credit interest", [("090", interest, 0), ("270", 0, interest)],
        "CREDIT INTEREST", "Interest income")
    jnl(pd, "Spend Money", f"FEE-{m:02d}", CO["bank"], "Account maintenance fee", [("404", 120, 0), ("090", 0, 120)],
        "ACCOUNT MAINTENANCE FEE", "Bank charges")
for d, vendor, acct, amt, desc, bd, cat in ONE_OFF:
    jnl(d, "Spend Money", f"SM-{d:%m%d}", vendor, desc, [(acct, float(amt), 0), ("090", 0, float(amt))], bd, cat)

JOURNALS.sort(key=lambda j: (j["date"], j["source"] != "Manual Journal", j["ref"]))
for n, j in enumerate(JOURNALS, 1):
    j["no"] = n
GL = [dict(date=j["date"], source=j["source"], no=j["no"], ref=j["ref"], contact=j["contact"], desc=j["desc"],
           code=a, dr=dr, cr=cr, bank_desc=j.get("bank_desc2") if (a, dr) == ("404", 150.0) or (a == "090" and cr == 150.0)
           else j["bank_desc"]) for j in JOURNALS for a, dr, cr in j["lines"]]


def balance(code, until):
    return r2(sum(r["dr"] - r["cr"] for r in GL if r["code"] == code and r["date"] <= until))


def actual(code, m):  # P&L sign: revenue positive as credit
    v = sum(r["dr"] - r["cr"] for r in GL if r["code"] == code and r["date"].month == m and r["date"].year == 2026)
    return r2(-v if ACC[code][2] == "Revenue" else v)


# ----------------------------------------------------------------------------- PDF engine
FONT_DIR = Path("/usr/share/fonts/noto-cjk")


def _idx(path):
    c = TTCollection(str(path), lazy=True)
    return next(i for i, f in enumerate(c.fonts) if f["name"].getDebugName(1).startswith("Noto") and
                f["name"].getDebugName(1).endswith(" CJK TC"))


FONTS = {("sans", ""): FONT_DIR / "NotoSansCJK-Regular.ttc", ("sans", "B"): FONT_DIR / "NotoSansCJK-Bold.ttc",
         ("serif", ""): FONT_DIR / "NotoSerifCJK-Regular.ttc", ("serif", "B"): FONT_DIR / "NotoSerifCJK-Bold.ttc"}
FONT_IDX = {k: _idx(p) for k, p in FONTS.items()}
PDF_DATE = datetime(2026, 10, 5, 9, 0, 0)


class Doc(FPDF):
    def __init__(self, footer=None, serif=False, when=PDF_DATE):
        super().__init__(format="A4")
        for (fam, st), p in FONTS.items():
            if fam == "sans" or serif:
                self.add_font(fam, st, str(p), collection_font_number=FONT_IDX[(fam, st)])
        self.footer_text = footer
        self.set_creation_date(when if isinstance(when, datetime) else datetime.combine(when, datetime.min.time()))
        self.set_margins(18, 15, 18)
        self.set_auto_page_break(True, 18)

    def footer(self):
        if self.footer_text:
            self.set_y(-13)
            self.set_font("sans", "", 7.5)
            self.set_text_color(110)
            self.cell(0, 5, self.footer_text.replace("{p}", str(self.page_no())), align="C")


def txt(p, s, size=10, style="", align="L", color=(0, 0, 0), fam="sans", h=None, w=0, gap=0):
    p.set_font(fam, style, size)
    p.set_text_color(*color)
    p.multi_cell(w, h or size * 0.48, s, align=align, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if gap:
        p.ln(gap)


def kv(p, x, y, items, wk, wv, size=9, h=5.2, bold_v=False):
    for i, (k, v) in enumerate(items):
        p.set_xy(x, y + i * h)
        p.set_font("sans", "", size)
        p.set_text_color(90)
        p.cell(wk, h, k)
        p.set_font("sans", "B" if bold_v else "", size)
        p.set_text_color(0)
        p.cell(wv, h, v)


def table(p, cols, rows, widths, aligns, size=8.5, fill=(232, 232, 232), h=6.2, border=1, text_color=(0, 0, 0)):
    p.set_font("sans", "B", size)
    p.set_fill_color(*fill)
    p.set_text_color(*text_color)
    for c, w in zip(cols, widths):
        p.cell(w, h, c, border=border, align="C", fill=True)
    p.ln(h)
    p.set_font("sans", "", size)
    p.set_text_color(0)
    for r in rows:
        for c, w, a in zip(r, widths, aligns):
            p.cell(w, h, str(c), border=border, align=a)
        p.ln(h)


def totals(p, items, x=112, wk=50, wv=30, h=6, size=9):
    for i, (k, v) in enumerate(items):
        last = i == len(items) - 1
        p.set_x(x)
        p.set_font("sans", "B" if last else "", size + (1 if last else 0))
        p.cell(wk, h, k, align="R")
        p.cell(wv, h, v, align="R", border="T" if last else 0)
        p.ln(h)


def stamp(p, x, y, top, bottom, color=(200, 30, 30), r=17):
    p.set_draw_color(*color)
    p.set_line_width(0.7)
    p.ellipse(x - r, y - r, 2 * r, 2 * r)
    p.set_text_color(*color)
    p.set_font("sans", "B", 7 if r >= 15 else 5.2)
    p.set_xy(x - r, y - 7)
    p.cell(2 * r, 4, top, align="C")
    p.set_font("sans", "B", 12)
    p.set_xy(x - r, y - 3)
    p.cell(2 * r, 6, "★", align="C")
    p.set_font("sans", "", 6.5)
    p.set_xy(x - r, y + 4)
    p.cell(2 * r, 4, bottom, align="C")
    p.set_draw_color(0)
    p.set_text_color(0)


def squiggle(p, x, y, seed, w=48):
    rr = random.Random(seed)
    pts = []
    for i in range(40):
        t = i / 39
        pts.append((x + t * w, y + 3.2 * math.sin(t * rr.uniform(9, 14)) * (1 - t * 0.5) + rr.uniform(-0.6, 0.6)))
    p.set_draw_color(20, 30, 110)
    p.set_line_width(0.45)
    p.polyline(pts)
    p.set_draw_color(0)
    p.set_line_width(0.2)


def rule(p, color=(0, 0, 0), wdt=0.4, gap=3):
    y = p.get_y() + 1
    p.set_draw_color(*color)
    p.set_line_width(wdt)
    p.line(p.l_margin, y, 210 - p.r_margin, y)
    p.set_line_width(0.2)
    p.set_draw_color(0)
    p.set_y(y + gap)


def out(p):
    return bytes(p.output())


def scan(pdf_bytes, seed):
    """Render to image, grayscale, rotate, add paper tone/noise/blur, embed as image-only PDF."""
    rr = random.Random(seed)
    page = fitz.open(stream=pdf_bytes, filetype="pdf")[0]
    pix = page.get_pixmap(dpi=150, colorspace=fitz.csGRAY)
    img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
    img = img.point(lambda v: int(v * 0.86 + 22))  # paper tone, lifted blacks
    img = img.rotate(rr.choice([-1, 1]) * rr.uniform(0.5, 1.5), resample=Image.BICUBIC, fillcolor=236)
    noise = Image.frombytes("L", img.size, rr.randbytes(img.size[0] * img.size[1]))
    img = Image.blend(img, noise, 0.07).filter(ImageFilter.GaussianBlur(0.6))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=62)
    p = FPDF(format="A4")
    p.set_creation_date(PDF_DATE)
    p.set_producer("DocuScan 3.2")
    p.add_page()
    p.image(io.BytesIO(buf.getvalue()), x=0, y=0, w=210, h=297)
    return out(p)



def clause(p, num, title, body, fam="sans", size=9.5, align=None):
    align = align or ("J" if fam == "serif" else "L")
    txt(p, f"{num}  {title}", size + 0.5, "B", fam=fam, gap=1)
    for sub, t in body:
        p.set_x(p.l_margin)
        p.set_font(fam, "", size)
        p.cell(10, size * 0.5, sub)
        p.multi_cell(0, size * 0.5, t, align=align, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        p.ln(1)
    p.ln(2)



# ----------------------------------------------------------------------------- document layouts
def inv_s01(inv):
    """Pearl River Growth Fund II drawdown notice (Chinese, mainland GP)."""
    s, red = SUP["S01"], (150, 25, 25)
    p = Doc(when=inv["date"])
    p.add_page()
    txt(p, s["name_zh"], 19, "B", "C", red)
    txt(p, s["name_en"].upper(), 10, "", "C", red)
    txt(p, f"普通合伙人：珠江资本管理（深圳）有限公司   {s['addr_zh']}   电话 {s['tel']}", 7.5, "", "C", (90, 90, 90))
    rule(p, red, 0.7, 4)
    txt(p, "缴 款 通 知 书", 16, "B", "C")
    txt(p, "CAPITAL CALL / DRAWDOWN NOTICE", 8, "", "C", (90, 90, 90), gap=4)
    y = p.get_y()
    kv(p, 18, y, [("有限合伙人：", CO["name_zh"]), ("", CO["name_en"]), ("", f"代表{CO['trust_zh']}"),
                  ("认购协议：", s["contract_no"])], 22, 86)
    kv(p, 126, y, [("通知编号：", inv["no"]), ("通知日期：", zh_date(inv["date"])), ("缴款截止日：", zh_date(inv["due"])),
                   ("币种：", "人民币 RMB")], 22, 44, bold_v=True)
    p.set_y(y + 26)
    commit_, before = s["commitment"], inv["called_before"]
    amt = inv["total"]
    rows = [("认缴出资额 Commitment", money(commit_)),
            ("此前累计实缴 Called before this notice", money(before)),
            (f"本次缴款比例 This call ({amt / commit_ * 100:.2f}% of commitment)", money(amt)),
            ("本次缴款后累计实缴 Called to date incl. this call", money(before + amt)),
            ("剩余未缴出资 Unfunded commitment after this call", money(commit_ - before - amt))]
    table(p, ["项目 Item", "金额（人民币元）"], rows, [124, 50], "LR", fill=(250, 228, 228))
    p.ln(3)
    txt(p, "资金用途：新项目投资及基金费用（详见附件一）。Use of proceeds: new portfolio investment and fund expenses.", 8.5, gap=2)
    totals(p, [("本次应缴金额 合计 人民币（RMB）", money(amt))], x=92, wk=70, wv=30)
    p.ln(6)
    txt(p, "收款账户信息 Wire instructions", 9.5, "B", color=red)
    y = p.get_y() + 1
    kv(p, 18, y, [("收款银行：", inv["bank_name"].split(" (")[1].rstrip(")")), ("账户名称：", s["name_zh"]),
                  ("银行账号：", inv["bank"])], 20, 100, bold_v=True)
    p.set_y(y + 20)
    txt(p, f"请于缴款截止日前电汇，并注明通知编号。如有疑问，请联系 {inv['email']}", 8, color=(90, 90, 90), gap=8)
    y = p.get_y()
    kv(p, 18, y, [("经办：", "张丽"), ("复核：", "陈国华")], 14, 30)
    stamp(p, 160, y + 6, "珠江资本管理有限公司", "基金专用章")
    return out(p)


def inv_s02(inv):
    """Harbourview Capital Partners III: capital call notice or quarterly management fee notice."""
    s, navy = SUP["S02"], (25, 45, 85)
    p = Doc(when=inv["date"])
    p.add_page()
    p.set_fill_color(*navy)
    p.rect(0, 0, 210, 28, "F")
    p.set_xy(18, 8)
    p.set_text_color(255)
    p.set_font("sans", "B", 13)
    p.cell(120, 7, "HARBOURVIEW CAPITAL PARTNERS III, L.P.")
    p.set_xy(18, 16)
    p.set_font("sans", "", 8.5)
    p.cell(120, 5, "Managed by Harbourview Capital Management Ltd")
    p.set_xy(130, 10)
    p.set_font("sans", "B", 11)
    p.cell(62, 7, "CAPITAL CALL NOTICE" if inv["kind"] == "capital_call" else "MANAGEMENT FEE NOTICE", align="R")
    p.set_y(33)
    txt(p, f"{s['addr_en']}  |  Tel {s['tel']}  |  {s['email']}", 7.5, color=(80, 80, 80), gap=1)
    txt(p, s["reg"], 7.5, color=(80, 80, 80), gap=5)
    y = p.get_y()
    kv(p, 18, y, [("Limited Partner", CO["name_en"]), ("", f"for the {CO['trust']}"),
                  *[("", a) for a in split2(CO["addr_en"], 44)]], 28, 76)
    kv(p, 120, y, [("Notice no.", inv["no"]), ("Notice date", en_date(inv["date"])), ("Payment due", en_date(inv["due"])),
                   ("LP account", "LP-0117"), ("Side letter", s["contract_no"])], 28, 44, bold_v=True)
    p.set_y(y + 32)
    if inv["kind"] == "capital_call":
        before, amt, c = inv["called_before"], inv["total"], s["commitment"]
        txt(p, f"Pursuant to clause 5.1 of the Limited Partnership Agreement, the General Partner calls capital from the "
               f"Limited Partner as follows:", 9, gap=2)
        rows = [("Commitment", money(c)), ("Contributed before this notice", money(before)),
                (f"This capital call ({amt / c * 100:.0f}% of commitment)", money(amt)),
                ("Contributed to date incl. this notice", money(before + amt)),
                ("Unfunded commitment after this notice", money(c - before - amt))]
        table(p, ["Capital account (USD)", "Amount"], rows, [124, 50], "LR", fill=(215, 225, 240))
    else:
        r = inv["rate"]
        txt(p, f"Management fee for the quarter {QTR[inv['period'][4:]][0]} - {QTR[inv['period'][4:]][1]} {inv['period'][:4]}, "
               f"payable quarterly in advance.", 9, gap=1)
        txt(p, f"Fee rate: {r:.2f}% p.a. on commitment of USD 10,000,000", 9, "B", gap=2)
        rows = [(c, d, f"{q:,}", money(u)) for c, d, q, u in inv["lines"]]
        table(p, ["Code", "Description", "Qty", "Amount USD"], rows, [24, 108, 12, 30], "LLRR", fill=(215, 225, 240), size=7.8)
    p.ln(3)
    totals(p, [("AMOUNT DUE (USD)", money(inv["total"]))])
    p.ln(6)
    txt(p, "Wire instructions", 9.5, "B", color=navy)
    txt(p, f"Beneficiary: {s['name_en']}\nBank: {s['bank']}\nAccount No.: {inv['bank']}\nReference: {inv['no']} / LP-0117",
        8.5, h=4.6, gap=4)
    txt(p, "Any change to these wire instructions will only ever be confirmed by a signed letter from the General Partner. "
           "Please call your investor relations contact on a known number before acting on any change.", 7.5, color=(90, 90, 90))
    return out(p)


def inv_s03(inv):
    s, green = SUP["S03"], (16, 110, 70)
    p = Doc(when=inv["date"])
    p.add_page()
    p.set_fill_color(*green)
    p.rect(18, 14, 22, 22, "F")
    p.set_xy(18, 20)
    p.set_text_color(255)
    p.set_font("sans", "B", 13)
    p.cell(22, 10, "PE", align="C")
    p.set_xy(60, 14)
    p.set_text_color(*green)
    p.set_font("sans", "B", 13)
    p.cell(132, 7, s["name_en"].upper(), align="R")
    p.set_text_color(80)
    p.set_font("sans", "", 8)
    for i, line in enumerate([f"{s['name_zh']}  ·  {s['addr_en']}", f"Tel {s['tel']}  |  {s['email']}", s["reg"]]):
        p.set_xy(44, 22 + i * 4.5)
        p.cell(148, 4.5, line, align="R")
    p.set_y(44)
    txt(p, "INVOICE", 22, "B", color=green, gap=3)
    y = p.get_y()
    kv(p, 18, y, [("Bill to", CO["name_en"]), *[("", x) for x in split2(CO["addr_en"], 50)]], 16, 90)
    kv(p, 130, y, [("Invoice no.", inv["no"]), ("Date", en_date(inv["date"])), ("Due date", en_date(inv["due"])),
                   ("Client", "LPFO-0042")], 22, 40, bold_v=True)
    p.set_y(y + 25)
    rows = [(c, d, q, PM_RATES[c][1], money(u), money(q * u)) for c, d, q, u in inv["lines"]]
    table(p, ["Code", "Description", "Qty", "Unit", "Rate HKD", "Amount HKD"], rows, [18, 86, 10, 14, 22, 24],
          "LLRCRR", fill=(210, 235, 222), border="B", size=7.8)
    p.ln(2)
    totals(p, [("Subtotal", money(inv["subtotal"])), ("TOTAL HKD", money(inv["total"]))])
    p.ln(8)
    txt(p, f"Payment within {s['terms']} days to {s['bank']}, account {inv['bank']} ({s['name_en']}).", 8.5)
    txt(p, "Rates per property management agreement 2026, schedule 1. E. & O. E.", 8, color=(100, 100, 100))
    return out(p)


def inv_s05(inv):
    s, purple = SUP["S05"], (92, 50, 140)
    p = Doc(when=inv["date"])
    p.add_page()
    p.set_text_color(*purple)
    p.set_font("sans", "B", 18)
    p.cell(110, 10, "Meridian Fine Art")
    p.set_font("sans", "", 18)
    p.set_text_color(150)
    p.cell(64, 10, "Invoice", align="R")
    p.ln(12)
    txt(p, f"{s['name_en']} · {s['addr_en']} · {s['reg']}", 7.5, color=(120, 120, 120), gap=5)
    y = p.get_y()
    kv(p, 18, y, [("Insured", CO["name_en"]), ("", f"Attn: {PEOPLE['grace'][0]}"), ("", PEOPLE["grace"][1])], 20, 80)
    kv(p, 128, y, [("Invoice #", inv["no"]), ("Issued", inv["date"].strftime("%B %d, %Y")),
                   ("Due", inv["due"].strftime("%B %d, %Y")), ("Policy", "MAI-FA-2025-0716")], 24, 40, bold_v=True)
    p.set_y(y + 26)
    rows = [(d, q, money(u), money(q * u)) for _, d, q, u in inv["lines"]]
    table(p, ["Description", "Qty", "Unit price", "Amount (USD)"], rows, [110, 10, 24, 30], "LRRR", border="B",
          fill=(236, 230, 250), size=7.8)
    p.ln(2)
    totals(p, [("Subtotal", money(inv["subtotal"])), ("Tax", "0.00"), ("Amount due (USD)", money(inv["total"]))])
    p.ln(8)
    txt(p, f"Pay by wire: {s['bank']}, account {inv['bank']}. Questions: {s['email']}", 8.5)
    return out(p)


RENDER = {"S01": inv_s01, "S02": inv_s02, "S03": inv_s03, "S05": inv_s05}


# ----------------------------------------------------------------------------- contracts
def signatures(p, left, right, when, stamp_zh=None):
    y = p.get_y()
    if y > 240:
        p.add_page()
        y = p.get_y()
    for x, (who, name) in [(18, left), (112, right)]:
        kv(p, x, y, [(who, ""), ("", ""), ("", ""), ("Name:", name), ("Date:", en_date(when))], 14, 60)
        squiggle(p, x + 4, y + 11, len(name) * 7 + x)
    if stamp_zh:
        stamp(p, 178, y - 2, stamp_zh, "合同专用章", r=14)


def subscription_agreement():
    """S01: subscription agreement excerpt (bilingual) with the designated account clause."""
    s = SUP["S01"]
    p = Doc(footer=f"{s['contract_no']}  ·  Page {{p}} / {{nb}}  ·  第{{p}}页", when=D(2024, 3, 18))
    p.add_page()
    txt(p, "SUBSCRIPTION AGREEMENT  认购协议", 16, "B", "C")
    txt(p, f"{s['name_en']}  {s['name_zh']}", 10, "", "C")
    txt(p, f"Agreement No. 协议编号: {s['contract_no']}   ·   Dated 签署日期: 18 March 2024", 9, "", "C", gap=5)
    txt(p, f"Limited Partner 有限合伙人: {CO['name_en']} {CO['name_zh']}, for the {CO['trust']}, BR No. {CO['br']}, {CO['addr_en']}", 9, gap=1)
    txt(p, f"Fund 基金: {s['name_en']} {s['name_zh']}, {s['reg']}, {s['addr_en']}", 9, gap=4)
    clause(p, "1", "Commitment 认缴出资", [
        ("1.1", f"The Limited Partner commits RMB {s['commitment']:,.0f} to the Fund. 有限合伙人认缴出资人民币{s['commitment'] / 1e4:,.0f}万元。"),
        ("1.2", "Commitments are drawn down by written drawdown notice from the General Partner. 出资按普通合伙人书面缴款通知分期实缴。")])
    clause(p, "2", "Term 期限", [("2.1", "Investment period 5 years from first close; fund term 8 years, extendable twice by one year. 投资期自首次交割起5年；基金存续期8年，可延长两次，每次一年。")])
    clause(p, "3", "Drawdowns 缴款", [
        ("3.1", "Each drawdown notice states the amount, the purpose and the due date, at least 10 business days after the notice. 每份缴款通知须列明金额、用途及缴款截止日，截止日不得早于通知日后10个工作日。"),
        ("3.2", "Late payment bears default interest at 8% per annum. 逾期缴款按年利率8%计收违约利息。")])
    clause(p, "4", "Fees 费用", [("4.1", "Management fee 2.0% per annum of commitment during the investment period, included in drawdowns. 投资期内管理费为认缴出资额的年2.0%，计入缴款。")])
    clause(p, "5", "Payment 付款", [
        ("5.1", "Drawdowns are paid by telegraphic transfer in RMB. 缴款以人民币电汇支付。"),
        ("5.2", f"The Fund's designated account 基金指定账户: {s['bank']}, account name {s['name_en']}, account no. {s['bank_acct']}."),
        ("5.3", "Any change of the designated account must be confirmed in writing by the General Partner's authorised signatories; "
                "email notice alone is not valid. 指定账户如有变更，须经普通合伙人授权代表书面确认，仅凭电子邮件通知无效。")])
    clause(p, "6", "General 一般条款", [
        ("6.1", "This Agreement is governed by the laws of the PRC; disputes go to the Shenzhen Court of International Arbitration.\n本协议适用中国法律，争议提交深圳国际仲裁院仲裁。"),
        ("6.2", "In case of conflict, the Chinese version prevails.\n中英文版本如有冲突，以中文版本为准。")])
    p.ln(4)
    signatures(p, ("For the Limited Partner 有限合伙人", "Victoria Cheung, Director"),
               ("For the General Partner 普通合伙人", "Zhang Wei, Managing Partner"), D(2024, 3, 18), "珠江资本管理有限公司")
    return out(p)


def side_letter():
    """S02: side letter granting the reduced management fee."""
    s = SUP["S02"]
    p = Doc(footer=f"Side letter {s['contract_no']}  ·  Page {{p}} of {{nb}}", serif=True, when=D(2024, 6, 3))
    p.add_page()
    txt(p, "HARBOURVIEW CAPITAL PARTNERS III, L.P.", 14, "B", "C", fam="serif")
    txt(p, "SIDE LETTER", 13, "B", "C", fam="serif")
    txt(p, f"Ref: {s['contract_no']}   ·   3 June 2024", 9, "", "C", (90, 90, 90), gap=6)
    txt(p, f"To: {CO['name_en']} (for the {CO['trust']}), {CO['addr_en']} (the \"Investor\")", 10, fam="serif", gap=4)
    txt(p, "Dear Investor,", 10, fam="serif", gap=2)
    txt(p, "In consideration of the Investor's commitment of USD 10,000,000 to Harbourview Capital Partners III, L.P. (the "
           "\"Fund\"), the General Partner agrees the following terms, which prevail over the Limited Partnership Agreement "
           "dated 15 May 2024 (the \"LPA\").", 10, fam="serif", gap=4)
    clause(p, "1.", "COMMITMENT", [("1.1", "The Investor's commitment is USD 10,000,000 (the \"Commitment\").")], "serif", 10)
    clause(p, "2.", "MOST FAVOURED NATION", [("2.1", "The Investor may elect the benefit of any more favourable fee terms granted to another investor with an equal or smaller commitment.")], "serif", 10)
    clause(p, "3.", "MANAGEMENT FEE", [
        ("3.1", f"Notwithstanding clause 8.1 of the LPA, the management fee payable by the Investor shall be calculated at the rate of "
                f"{FEE_RATE:.2f}% per annum of the Commitment (instead of 2.00% per annum) for the whole term of the Fund."),
        ("3.2", "The management fee is payable quarterly in advance on the first business day of each calendar quarter."),
        ("3.3", "No other fee rate applies to the Investor unless agreed in a further side letter signed by both parties.")], "serif", 10)
    clause(p, "4.", "REPORTING", [("4.1", "Quarterly capital account statements within 60 days of each quarter end.")], "serif", 10)
    clause(p, "5.", "WIRE INSTRUCTIONS", [("5.1", f"Capital calls and fees are paid to {s['bank']}, account {s['bank_acct']}. "
                                                  "Changes are valid only by signed letter from the General Partner.")], "serif", 10)
    p.ln(4)
    signatures(p, ("For the General Partner", "Michael Tan, Managing Partner"), ("Agreed by the Investor", "Victoria Cheung, Director"),
               D(2024, 6, 3))
    return out(p)


def pm_agreement():
    s, green = SUP["S03"], (16, 110, 70)
    p = Doc(when=D(2025, 12, 1))
    p.add_page()
    txt(p, s["name_en"].upper(), 15, "B", color=green)
    txt(p, f"{s['addr_en']} · {s['reg']}", 8, color=(90, 90, 90), gap=4)
    txt(p, "PROPERTY MANAGEMENT AGREEMENT 2026", 14, "B", gap=1)
    txt(p, f"Client: {CO['name_en']} for the {CO['trust']} (client LPFO-0042)   ·   Term 1 Jan - 31 Dec 2026", 9, gap=1)
    txt(p, f"Properties: {PROPS['12A']}; {PROPS['21B']}", 8.5, gap=4)
    txt(p, "Schedule 1 - Fees", 10, "B", gap=1)
    table(p, ["Code", "Service", "Unit", "Rate (HKD)"], [(c, d, u, money(r)) for c, (d, u, r) in PM_RATES.items()],
          [20, 110, 18, 26], "LLCR", fill=(210, 235, 222), size=8)
    p.ln(4)
    for n in ["Repairs above HK$5,000 need the client's prior written approval; contractors are paid by the client directly.",
              "Leasing commission on new tenancies or renewals: half of one month's rent, payable on signing.",
              f"Invoices are issued monthly; payment terms {s['terms']} days."]:
        txt(p, "•  " + n, 9, gap=1)
    p.ln(6)
    txt(p, f"Accepted for the client: {PEOPLE['grace'][0]}, Fund Accountant, 3 Dec 2025", 9)
    squiggle(p, 30, p.get_y() + 6, 77)
    return out(p)


LEASE_NEW = dict(ref="LPFO/L/2027/12A", signed=D(2026, 9, 29), start=D(2027, 1, 1), end=D(2029, 12, 31), rent=98800.0,
                 esc=0.03, deposit=296400, words="Ninety-Eight Thousand Eight Hundred")
LEASE_OLD = dict(ref="LPFO/L/2024/12A", signed=D(2023, 11, 20), start=D(2024, 1, 1), end=D(2026, 12, 31), rent=95000.0,
                 esc=0.0, deposit=285000, words="Ninety-Five Thousand")


def rent_schedule(L):
    return [(L["start"].year + i, r2(L["rent"] * (1 + L["esc"]) ** i)) for i in range(L["end"].year - L["start"].year + 1)]


BOILER_LEASE = [
    "The Tenant shall keep the interior of the Premises and the Landlord's fixtures and fittings in good and tenantable repair.",
    "The Tenant shall not assign, sublet or part with possession of the Premises.",
    "The Tenant shall use the Premises as a private residence for its nominated executive and family only.",
    "The Landlord shall keep the main structure, roof and exterior walls in proper repair.",
    "The Tenant shall maintain contents and public liability insurance of at least HK$5,000,000.",
    "No alterations shall be made to the Premises without the Landlord's prior written consent.",
    "Notices may be served at the addresses stated above by hand or by registered post.",
    "Time shall be of the essence in respect of all payments under this Agreement.",
]


def lease(L):
    s = SUP["S04"]
    new = L["esc"] > 0
    p = Doc(footer=f"Tenancy Agreement {L['ref']}  ·  Page {{p}} of {{nb}}", serif=True, when=L["signed"])
    p.add_page()
    txt(p, "TENANCY AGREEMENT", 17, "B", "C", fam="serif")
    txt(p, "租 約", 12, "", "C", fam="serif")
    txt(p, f"Ref: {L['ref']}", 9, "", "C", (90, 90, 90), gap=6)
    txt(p, f"THIS AGREEMENT is made on {en_date(L['signed'])}", 10.5, fam="serif", gap=3)
    txt(p, "BETWEEN", 10.5, "B", fam="serif", gap=2)
    txt(p, f"(1)  {CO['name_en'].upper()} ({CO['name_zh']}), BR No. {CO['br']}, for and on behalf of the trustee of the "
           f"{CO['trust']}, whose registered office is at {CO['addr_en']} (the \"Landlord\"); and", 10.5, fam="serif", gap=2)
    txt(p, f"(2)  {s['name_en'].upper()}, {s['reg']}, whose registered office is at {s['addr_en']} (the \"Tenant\").", 10.5, fam="serif", gap=5)
    txt(p, "IT IS AGREED as follows:", 10.5, fam="serif", gap=4)
    clause(p, "1.", "PREMISES", [("1.1", f"The Landlord lets and the Tenant takes {PROPS['12A']} with a saleable area of approximately 2,150 sq. ft. and one car parking space (the \"Premises\")."),
                                 ("1.2", "The Premises shall be used as a private residence only.")], "serif", 10.5)
    if new:
        p.add_page()
    clause(p, "2.", "TERM", [("2.1", f"The term is {L['end'].year - L['start'].year + 1} years from {en_date(L['start'])} to {en_date(L['end'])}, both days inclusive (the \"Term\").")], "serif", 10.5)
    clause(p, "3.", "DEPOSIT AND TENANT'S OBLIGATIONS",
           [("3.1", f"The Tenant shall pay a security deposit of HK${L['deposit']:,} on signing, refundable without interest on expiry.")] +
           [(f"3.{i}", t) for i, t in enumerate(BOILER_LEASE if new else BOILER_LEASE[:4], 2)], "serif", 10.5)
    p.add_page()
    if new:
        body = [("4.1", f"The monthly rent shall be HK${L['rent']:,.0f} (Hong Kong Dollars {L['words']}) "
                        f"inclusive of management fees and exclusive of rates, payable from {en_date(L['start'])}."),
                ("4.2", "On each 1 January during the Term the monthly rent shall increase by three per cent (3%) over the monthly rent payable in the preceding year."),
                ("4.3", "For the avoidance of doubt, the monthly rent for each year of the Term is:")]
    else:
        body = [("4.1", f"The monthly rent shall be HK${L['rent']:,.0f} (Hong Kong Dollars {L['words']}) inclusive of management fees and exclusive of rates, fixed for the whole Term."),
                ("4.2", "No rent review shall take place during the Term.")]
    clause(p, "4.", "RENT" + (" AND RENT REVIEW" if new else ""), body, "serif", 10.5)
    if new:
        table(p, ["Period", "Monthly rent (HK$)"], [(f"1 Jan {y} - 31 Dec {y}", money(r)) for y, r in rent_schedule(L)],
              [70, 45], "LR", size=9.5)
        p.ln(3)
    clause(p, "5.", "PAYMENT AND RATES",
           [("5.1", f"Rent is payable monthly in advance on the first day of each month by transfer to the Landlord's account with {CO['bank']}, a/c {CO['bank_acct']}."),
            ("5.2", "Government rates are payable by the Tenant.")], "serif", 10.5)
    clause(p, "6.", "GOVERNING LAW", [("6.1", "This Agreement is governed by the laws of the Hong Kong Special Administrative Region.")], "serif", 10.5)
    p.ln(4)
    txt(p, "IN WITNESS whereof the parties have signed this Agreement on the date first written above.", 10, fam="serif", gap=8)
    y = p.get_y()
    for x, who, name, org in [(18, "SIGNED by the Landlord", "Victoria Cheung, Director", CO["name_en"]),
                              (112, "SIGNED by the Tenant", "Daniel Fong, Managing Director", s["name_en"])]:
        p.set_xy(x, y)
        p.set_font("serif", "B", 9.5)
        p.cell(80, 5, who)
        p.set_xy(x, y + 22)
        p.set_draw_color(0)
        p.line(x, y + 21, x + 70, y + 21)
        p.set_font("serif", "", 9)
        p.cell(80, 5, f"{name}, for and on behalf of")
        p.set_xy(x, y + 27)
        p.cell(80, 5, org)
        squiggle(p, x + 6, y + 15, x * 3 + L["signed"].year)
    stamp(p, 98, y + 9, CO["name_zh"], "LANDLORD", (190, 30, 30), 12)
    p.set_y(y + 40)
    txt(p, f"Witness: {PEOPLE['grace'][0]}, Fund Accountant, {CO['name_en']}", 9, fam="serif")
    return out(p)


def insurance_policy():
    s, purple = SUP["S05"], (92, 50, 140)
    p = Doc(when=D(2025, 11, 10))
    p.add_page()
    txt(p, "Meridian Fine Art", 20, "B", color=purple)
    txt(p, "FINE ART COLLECTION POLICY - SCHEDULE   MAI-FA-2025-0716", 12, "B", gap=1)
    txt(p, f"{s['name_en']}, {s['addr_en']}", 8, color=(110, 110, 110), gap=5)
    y = p.get_y()
    kv(p, 18, y, [("Insured", f"{CO['name_en']} for the {CO['trust']}"), ("Billing contact", PEOPLE["grace"][1]),
                  ("Policy start", "November 16, 2025"), ("Policy period", "12 months (Nov 16, 2025 - Nov 15, 2026)"),
                  ("Renewal date", "November 16, 2026"), ("Premium", "USD 29,400 a year, 12 monthly instalments, net 15")], 34, 130)
    p.set_y(y + 36)
    table(p, ["Cover", "Sum insured", "Basis", "Monthly instalment"],
          [("Paintings, ceramics and sculpture (schedule of 64 works)", "USD 42,000,000", "Agreed value", "USD 2,450.00")],
          [84, 32, 26, 32], "LRCR", fill=(236, 230, 250), size=8)
    p.ln(5)
    terms = [("1. Auto-renewal.", "This policy renews automatically on the Renewal Date for a further 12 months at the insurer's then-current "
                                  "rates unless the Insured gives written notice of non-renewal at least 30 days before the Renewal Date."),
             ("2. Notice.", "Non-renewal notices must be sent in writing to underwriting@meridianfineart.com. Notices via the broker portal chat are not valid."),
             ("3. Premium.", "Instalments are non-refundable once the policy renews. Works may be added by endorsement at any time."),
             ("4. Conditions.", "Works must be kept at the declared locations: the family residence and the Lantau Peak art store, Kwai Chung.")]
    for h, t in terms:
        txt(p, h, 9.5, "B")
        txt(p, t, 9, gap=2)
    p.ln(6)
    txt(p, f"Signed for the Insured: {PEOPLE['grace'][0]}, Fund Accountant — November 10, 2025", 9)
    squiggle(p, 30, p.get_y() + 6, 31)
    return out(p)


def bank_statement(lines, opening):
    navy = (16, 40, 90)
    p = Doc(footer="Victoria Harbour Bank Limited (fictional) · Statement 2026-09 · Page {p} of {nb}", when=D(2026, 10, 2))
    p.add_page()
    p.set_fill_color(*navy)
    p.rect(0, 0, 210, 24, "F")
    p.set_xy(18, 7)
    p.set_text_color(255)
    p.set_font("sans", "B", 15)
    p.cell(120, 8, "VICTORIA HARBOUR BANK")
    p.set_font("sans", "", 9)
    p.set_xy(120, 8)
    p.cell(72, 6, "維港銀行  ·  Private Client Current Account", align="R")
    p.set_y(30)
    kv(p, 18, 30, [("Account name", CO["name_en"]), ("Account no.", CO["bank_acct"] + " (HKD Current)"),
                   ("Statement period", "01 Sep 2026 - 30 Sep 2026"), ("Branch", "Central Private Client Centre")], 30, 100)
    p.set_y(54)
    rows, bal = [], opening
    rows.append(("01 Sep", "BALANCE BROUGHT FORWARD", "", "", money(opening)))
    for d, desc, amt in lines:
        bal = r2(bal + amt)
        rows.append((d.strftime("%d %b"), desc[:60], money(-amt) if amt < 0 else "", money(amt) if amt > 0 else "", money(bal)))
    rows.append(("30 Sep", "CLOSING BALANCE", "", "", money(bal)))
    table(p, ["Date", "Transaction details", "Withdrawal", "Deposit", "Balance"], rows, [14, 86, 24, 24, 28],
          "LLRRR", size=7.2, h=5.2, border="B", fill=(220, 226, 240))
    p.ln(4)
    dep, wd = sum(a for _, _, a in lines if a > 0), -sum(a for _, _, a in lines if a < 0)
    totals(p, [("Opening balance", money(opening)), ("Total deposits", money(dep)), ("Total withdrawals", money(wd)),
               ("Closing balance HKD", money(bal))], size=8.5)
    return out(p), bal


# ----------------------------------------------------------------------------- sheets
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
NUM = "#,##0.00"


def header(ws, cols, widths, row=1):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font, cell.alignment = HDR_FILL, HDR_FONT, Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = f"A{row + 1}"


def save_wb(wb, rel, when):
    wb.properties.creator = "Grace Lam"
    wb.properties.lastModifiedBy = "Jason Yip"
    stamp_ = when.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    buf = io.BytesIO()
    wb.save(buf)
    zin, ob_ = zipfile.ZipFile(buf), io.BytesIO()
    with zipfile.ZipFile(ob_, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            b = zin.read(it.filename)
            if it.filename == "docProps/core.xml":
                b = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>" + stamp_, b)
            zi = zipfile.ZipInfo(it.filename, date_time=when.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, b)
    write(rel, ob_.getvalue())


def reg_rows(until, as_of):
    rows = []
    for inv in sorted([i for i in HIST if i["date"] <= until], key=lambda i: (i["date"], i["sid"], i["no"])):
        s = SUP[inv["sid"]]
        status = "Paid" if inv["pay_date"] <= as_of else "Approved"
        rows.append([inv["date"], inv["sid"], s["name_en"], inv["no"], inv["currency"], inv["total"],
                     fx(inv["currency"], inv["date"]), inv["amount_hkd"], inv["account"], inv["due"], status,
                     f"{inv['folder']}/{inv['file']}"])
    return rows


REG_COLS = ["date", "supplier_id", "supplier", "invoice_no", "currency", "amount", "fx_rate", "amount_hkd",
            "account_code", "due_date", "status", "source_file"]


def write_register(until, as_of, when):
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    header(ws, REG_COLS, [11, 10, 36, 18, 9, 16, 8, 16, 12, 11, 10, 56])
    for r, row in enumerate(reg_rows(until, as_of), 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=8, value=f"=ROUND(F{r}*G{r},2)")
        for c in (1, 10):
            ws.cell(row=r, column=c).number_format = "yyyy-mm-dd"
        for c in (6, 8):
            ws.cell(row=r, column=c).number_format = NUM
        ws.cell(row=r, column=7).number_format = "0.0000"
    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    save_wb(wb, "sheets/invoice_register.xlsx", when)


def write_master(when):
    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    cols = ["id", "name_en", "name_zh", "email_domain", "bank_name", "bank_account", "currency", "payment_terms", "contract_file"]
    header(ws, cols, [6, 38, 20, 22, 60, 22, 9, 14, 50])
    for sid, s in SUP.items():
        ws.append([sid, s["name_en"], s["name_zh"], s["domain"], s["bank"], s["bank_acct"], s["currency"],
                   f"{s['terms']} days", s["contract"]])
    save_wb(wb, "sheets/supplier_master.xlsx", when)


OTHER_FUNDS = [  # paid from the custodian account, not this bank account: no documents in the demo
    ("Kowloon Ventures Fund I LP", "", "USD", 3_000_000.0, 2_850_000.0, D(2025, 6, 12), "Custodian statement Jun 2025"),
    ("Asia Credit Opportunities Fund LP", "", "USD", 5_000_000.0, 3_250_000.0, D(2026, 4, 20), "Custodian statement Apr 2026"),
]


def commitments(until):
    rows = []
    for sid in ("S01", "S02"):
        s = SUP[sid]
        calls = [i for i in INVOICES if i["sid"] == sid and i["kind"] == "capital_call" and i["date"] <= until]
        last = max(calls, key=lambda i: i["date"]) if calls else None
        src = (f"{last['folder']}/{last['file']}" if last["folder"] else f"Notice {last['no']}") if last else "LP capital statement Dec 2025"
        rows.append([s["name_en"], sid, s["currency"], s["commitment"], s["called_before"] + sum(i["total"] for i in calls),
                     None, last["date"] if last else None, src])
    for name, sid, ccy, c, called, last, src in OTHER_FUNDS:
        rows.append([name, sid, ccy, c, called, None, last, src])
    return rows


def write_commitments(until, when):
    wb = Workbook()
    ws = wb.active
    ws.title = "Commitments"
    header(ws, ["fund", "counterparty_id", "currency", "commitment", "called_to_date", "unfunded", "last_call_date", "source"],
           [38, 14, 9, 16, 16, 16, 14, 58])
    for r, row in enumerate(commitments(until), 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=6, value=f"=D{r}-E{r}")
        for c in (4, 5, 6):
            ws.cell(row=r, column=c).number_format = NUM
        ws.cell(row=r, column=7).number_format = "yyyy-mm-dd"
    n = ws.max_row
    ws.cell(row=n + 2, column=1, value=f"Called to date as of {en_date(until)}. Funds without a counterparty_id are paid from the custodian account.")
    save_wb(wb, "sheets/commitments.xlsx", when)


def write_fx(until_month, when):
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly"
    header(ws, ["month", "USD_HKD", "RMB_HKD", "source"], [10, 10, 10, 40])
    for k, (u, r) in FX.items():
        if k <= until_month:
            ws.append([k, u, r, "Month-start rate, Victoria Harbour Bank board rate"])
    save_wb(wb, "sheets/fx_rates.xlsx", when)


def write_bva(until_m, when):
    wb = Workbook()
    ws = wb.active
    ws.title = "YTD"
    ws["A1"] = f"{CO['name_en']} - Budget vs Actual 2026, YTD {'Jan-' + D(2026, until_m, 1).strftime('%b') if until_m else '(no actuals yet)'}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Budget approved by V. Cheung 20 Jan 2026. Actuals from GL export. Variance = Actual - Budget (expenses: positive = over budget)."
    header(ws, ["code", "account", "group", "budget_ytd", "actual_ytd", "variance", "variance_pct"], [7, 38, 20, 14, 14, 14, 12], row=4)
    r = 5
    for code in PL_CODES:
        b = sum(budget(code, m) for m in range(1, until_m + 1)) if until_m else sum(budget(code, m) for m in range(1, 13))
        a = sum(actual(code, m) for m in range(1, until_m + 1))
        ws.append([code, ACC[code][1], ACC[code][3], b, r2(a), f"=E{r}-D{r}", f'=IF(D{r}=0,"",F{r}/D{r})'])
        r += 1
    last = r - 1
    ws.append([])
    r += 1
    ws.cell(row=r, column=2, value="By group").font = Font(bold=True)
    for g in GROUPS:
        r += 1
        ws.append(["", g, "group total", f'=SUMIF($C$5:$C${last},B{r},D$5:D${last})', f'=SUMIF($C$5:$C${last},B{r},E$5:E${last})',
                   f"=E{r}-D{r}", f'=IF(D{r}=0,"",F{r}/D{r})'])
    for row in ws.iter_rows(min_row=5, min_col=4, max_col=7):
        for c in row:
            c.number_format = "0.0%" if c.column == 7 else "#,##0"
    if not until_m:
        ws.cell(row=4, column=4, value="budget_fy")
    m_ws = wb.create_sheet("Budget monthly")
    header(m_ws, ["code", "account"] + [D(2026, m, 1).strftime("%b") for m in range(1, 13)], [7, 38] + [11] * 12)
    for code in PL_CODES:
        m_ws.append([code, ACC[code][1]] + [budget(code, m) for m in range(1, 13)])
    a_ws = wb.create_sheet("Actual monthly")
    header(a_ws, ["code", "account"] + [D(2026, m, 1).strftime("%b") for m in range(1, until_m + 1)], [7, 38] + [12] * 9)
    for code in PL_CODES:
        a_ws.append([code, ACC[code][1]] + [actual(code, m) for m in range(1, until_m + 1)])
    save_wb(wb, "sheets/budget_vs_actual_2026.xlsx", when)


FC_ROWS = [  # label, account(s), 2027 monthly, 2028 growth
    ("Rental income - Flat 12A Repulse Bay", "260", 95_000, 0.0),
    ("Rental income - Flat 21B Mid-Levels", "260", 62_000, 0.0),
    ("Portfolio income (dividends & interest)", "270/275", 3_950_000, 0.03),
    ("Total income", "", None, None),
    ("Staff (office & household)", "477/478", 455_000, 0.04),
    ("Property costs", "471/473/475", 34_000, 0.03),
    ("Education", "481", 52_000, 0.05),
    ("Travel & aviation", "493", 130_000, 0.03),
    ("Insurance", "433", 23_500, 0.03),
    ("Professional fees", "412", 38_000, 0.03),
    ("Fund management fees", "455", 97_500, 0.0),
    ("Other office costs", "4xx", 15_500, 0.03),
    ("Total expenses", "", None, None),
    ("Net cash from operations", "", None, None),
    ("Expected fund capital calls", "150", 2_600_000, -0.30),
    ("Net cash after capital calls", "", None, None),
]
RENT_ROW = 5
RENT_RANGE = "RENT_INCOME_FORECAST"


def write_forecast(when):
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws["A1"] = f"{CO['name_en']} - Family office cash forecast 2027-2028 (HKD)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "v1, 14 Aug 2026, prepared by G. Lam. Repulse Bay rent flat pending lease renewal (R. Ho)."
    mons = [D(2027 + i // 12, i % 12 + 1, 1) for i in range(24)]
    header(ws, ["line", "acct"] + [m.strftime("%b-%y") for m in mons], [38, 11] + [12] * 24, row=4)
    for i, (label, acct, base, g) in enumerate(FC_ROWS):
        r = 5 + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=acct)
        for j in range(24):
            col = get_column_letter(3 + j)
            formulas = {"Total income": f"=SUM({col}5:{col}7)", "Total expenses": f"=SUM({col}9:{col}16)",
                        "Net cash from operations": f"={col}8-{col}17", "Net cash after capital calls": f"={col}18-{col}19"}
            v = formulas.get(label) or round(base * (1 + g) ** (j // 12), -2)
            ws.cell(row=r, column=3 + j, value=v).number_format = "#,##0"
        if base is None:
            for c in ws[r]:
                c.font = Font(bold=True)
    assert ws.cell(row=RENT_ROW, column=1).value.startswith("Rental income - Flat 12A")
    wb.defined_names[RENT_RANGE] = DefinedName(RENT_RANGE, attr_text=f"Forecast!$C${RENT_ROW}:$Z${RENT_ROW}")
    a = wb.create_sheet("Assumptions")
    header(a, ["assumption", "value", "source"], [40, 30, 64])
    for row in [("Portfolio income 2028", "+3%", "R. Ho email 14 Aug 2026"),
                ("Headcount", "4 office + 6 household, flat", "R. Ho email 14 Aug 2026"),
                ("Rental income - Flat 12A Repulse Bay 2027-2028", "HK$95,000/month flat",
                 "Pending lease renewal; COMPANY.md; R. Ho email 14 Aug 2026"),
                ("FX", "USD 7.80, RMB 1.08", "fx_rates.xlsx"),
                ("Capital calls", "HK$2.6m/month 2027, -30% 2028", "commitments.xlsx; GP pacing guidance")]:
        a.append(row)
    save_wb(wb, "sheets/forecast_2027_2028.xlsx", when)


# ----------------------------------------------------------------------------- ledger CSVs, text docs
def write_csv(rel, cols, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(cols)
    w.writerows(rows)
    write(rel, buf.getvalue().encode())


def write_coa():
    write_csv("ledger/chart_of_accounts.csv", ["Code", "Name", "Type", "Report group", "Tax"],
              [(c, n, t, g, "Tax Exempt") for c, n, t, g in sorted(ACCOUNTS)])


def write_gl(until):
    write_csv("ledger/gl_export_2026.csv",
              ["Date", "Source", "Journal No.", "Reference", "Contact", "Description", "Account Code", "Account",
               "Debit", "Credit", "Net", "Tax Rate"],
              [(r["date"].isoformat(), r["source"], r["no"], r["ref"], r["contact"], r["desc"], r["code"], ACC[r["code"]][1],
                f"{r['dr']:.2f}" if r["dr"] else "", f"{r['cr']:.2f}" if r["cr"] else "", f"{r['dr'] - r['cr']:.2f}", "Tax Exempt")
               for r in GL if r["date"] <= until])


def write_payments(until):
    cols = list(PAYMENTS[0])
    write_csv("ledger/payments_history.csv", cols,
              [[p[c].isoformat() if isinstance(p[c], date) else (f"{p[c]:.2f}" if isinstance(p[c], float) else p[c])
                for c in cols] for p in sorted(PAYMENTS, key=lambda x: (x["payment_date"], x["supplier_id"])) if p["payment_date"] <= until])


def eml(rel, frm, to, subject, when, body, cc=None, extra=None):
    m = EmailMessage(policy=SMTPUTF8)
    m["From"], m["To"], m["Subject"] = frm, to, subject
    if cc:
        m["Cc"] = cc
    m["Date"] = when.strftime("%a, %d %b %Y %H:%M:%S +0800")
    m["Message-ID"] = f"<{hashlib.md5(subject.encode()).hexdigest()[:16]}@{frm.split('@')[1].rstrip('>')}>"
    for k, v in (extra or {}).items():
        m[k] = v
    m.set_content(body, cte="8bit")
    write(rel, m.as_bytes())


RENT_ASSUMPTION = "Rental income for Flat 12A Repulse Bay (Halcyon Re lease) assumed flat at HK$95,000/month in 2027-2028 pending lease renewal (R. Ho, Aug 2026)."


def company_md(with_forecast):
    sup = "\n".join(f"| {k} | {s['name_en']} {s['name_zh']} | {s['currency']} | {s['domain']} | {s['terms']} days | {s['contract']} |"
                    for k, s in SUP.items())
    fc = (f"\n## Forecast assumptions (2027-2028)\n- {RENT_ASSUMPTION}\n"
          "- Forecast rent cells: named range RENT_INCOME_FORECAST in sheets/forecast_2027_2028.xlsx.\n"
          "- Portfolio income +3% in 2028; staff flat (4 office + 6 household); FX USD 7.80, RMB 1.08.\n"
          "- Source: docs/emails/2026-08-14_raymond_forecast_assumptions.eml\n"
          if with_forecast else "")
    return f"""# COMPANY.md - {CO['name_en']} ({CO['name_zh']})

Hong Kong single-family office for the {CO['family']}. Manages about US$350M for the {CO['trust']}: a listed
portfolio at Lion Rock Private Bank (custodian), four private fund commitments, two rental flats in Hong Kong,
an art collection and the family's expenses (staff, school fees, travel and aviation, insurance).
Functional currency HKD (also USD, RMB). FY = calendar year. BR No. {CO['br']}.
Office: {CO['addr_en']}.
Bank: {CO['bank']} a/c {CO['bank_acct']} (operating account; fund calls, fees and family bills are paid from here).
Email domain: {CO['domain']}

## People
- Victoria Cheung - Principal (family member). Approves anything above HK${APPROVAL_LIMIT:,}.
- Raymond Ho - CFO. Forecast, budget, investment reporting.
- Grace Lam - Fund Accountant. Month-end close, fund capital accounts, invoice register.
- Jason Yip - Accounts Clerk. Document intake, payments run.

## Conventions
- Account codes (ledger/chart_of_accounts.csv): capital calls -> 150 (Investments - Private Funds);
  fund management fee notices -> 455 (Fund management fees); 260 Rental income; 433 Insurance; 471 Property management;
  473 Repairs; 475 Building management & rates. Property = 471 + 473 + 475.
- Invoice register: sheets/invoice_register.xlsx, one row per payable document (capital call, fee notice, invoice),
  amount_hkd = amount x month fx rate.
- Fund commitments: sheets/commitments.xlsx (commitment, called to date, unfunded).
- Payable PDFs filed under docs/invoices/YYYY-MM/; new ones arrive in docs/invoices/inbox/.

## Counterparties
| ID | Name | Ccy | Known domain | Terms | Contract |
|---|---|---|---|---|---|
{sup}

## Policies
- Price tolerance: an invoice unit price or fee rate may not exceed the contract or side-letter rate by more than 1%. Otherwise hold.
- Approval: any payment above HK${APPROVAL_LIMIT:,} needs V. Cheung's approval.
- Payments only to bank accounts on the counterparty master. Wire-instruction changes need a signed letter from the
  counterparty's authorised signatory plus a call-back to a number on file. Email alone is never enough.
- Duplicates: never enter a document whose number or counterparty+amount matches an existing entry.
{fc}"""


# ----------------------------------------------------------------------------- file/git plumbing
def write(rel, data):
    p = WS / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data if isinstance(data, bytes) else data.encode())


def git(*args, when=None, who="grace"):
    env = dict(os.environ, GIT_CONFIG_GLOBAL="/dev/null", GIT_CONFIG_SYSTEM="/dev/null",
               GIT_DIR=str(GIT_DIR), GIT_WORK_TREE=str(WS))
    if when:
        name, mail, _ = PEOPLE[who]
        env.update(GIT_AUTHOR_NAME=name, GIT_AUTHOR_EMAIL=mail, GIT_COMMITTER_NAME=name, GIT_COMMITTER_EMAIL=mail,
                   GIT_AUTHOR_DATE=when, GIT_COMMITTER_DATE=when)
    return subprocess.run(["git", *args], cwd=WS, env=env, check=True, capture_output=True, text=True).stdout.strip()


def commit(msg, when, who):
    git("add", "-A")
    git("commit", "-q", "-m", msg, when=when, who=who)
    return git("rev-parse", "HEAD")


def text_of(data, page=None):
    doc = fitz.open(stream=data, filetype="pdf")
    return "\n".join(doc[i].get_text() for i in (range(len(doc)) if page is None else [page]))


def find_page(data, needle):
    doc = fitz.open(stream=data, filetype="pdf")
    return next(i + 1 for i in range(len(doc)) if needle in doc[i].get_text().replace("\n", " "))


# ----------------------------------------------------------------------------- main
def main():
    if WS.exists():
        shutil.rmtree(WS)
    GIT_DIR.mkdir(parents=True)
    git("init", "-q", "-b", "main")
    git("config", "core.worktree", "../..")
    (GIT_DIR / "info").mkdir(exist_ok=True)
    (GIT_DIR / "info" / "exclude").write_text("/.trace/git/\n")

    # ---- render documents once
    pdfs = {}
    for inv in HIST + INBOX:
        b = RENDER[inv["sid"]](inv)
        pdfs[f"{inv['folder']}/{inv['file']}"] = scan(b, inv["no"]) if inv["file"] in SCANNED else b
        inv["clean_pdf"] = b
    contracts = {"docs/contracts/S01_subscription_agreement_2024.pdf": subscription_agreement(),
                 "docs/contracts/S02_side_letter_2024.pdf": side_letter(),
                 "docs/contracts/S03_property_management_agreement_2026.pdf": pm_agreement(),
                 "docs/contracts/S04_lease_2024_2026.pdf": lease(LEASE_OLD),
                 "docs/contracts/S05_art_insurance_policy_2025.pdf": insurance_policy()}
    new_lease = lease(LEASE_NEW)
    sep_bank = [r for r in GL if r["code"] == "090" and r["date"].month == 9]
    opening = balance("090", D(2026, 8, 31))
    stmt_lines = [(r["date"], r["bank_desc"] or r["desc"].upper(), r2(r["dr"] - r["cr"])) for r in sep_bank]
    stmt_pdf, stmt_close = bank_statement(stmt_lines, opening)

    # ---- assertions
    for j in JOURNALS:
        assert abs(sum(l[1] - l[2] for l in j["lines"])) < 0.005, j
    assert abs(sum(r["dr"] - r["cr"] for r in GL)) < 0.005
    run = OPEN_CASH
    for r in sorted([r for r in GL if r["code"] == "090"], key=lambda r: (r["date"], r["no"])):
        run += r["dr"] - r["cr"]
        assert run > 0, "cash went negative"
    assert abs(stmt_close - balance("090", CLOSE)) < 0.005, "bank statement does not reconcile to GL"
    stmt_set = {(d, -a) for d, _, a in stmt_lines}
    for pmt in PAYMENTS:
        if pmt["payment_date"].month == 9 and pmt["payment_date"].year == 2026:
            assert (pmt["payment_date"], pmt["amount_hkd"]) in stmt_set, pmt
    s01 = [p for p in PAYMENTS if p["supplier_id"] == "S01"]
    assert len(s01) >= 6 and all(p["bank_account_paid"].endswith("2049") for p in s01)
    for row in reg_rows(CLOSE, CLOSE):
        inv = by_no[row[3]]
        assert row[5] == inv["total"] and money(inv["total"]) in text_of(inv["clean_pdf"]), row
        assert abs(r2(row[5] * row[6]) - row[7]) < 0.005
        assert inv["date"].weekday() < 5
        assert abs(inv["total"] - sum(q * u for _, _, q, u in inv["lines"])) < 0.005
    for inv in INBOX:
        assert money(inv["total"]) in text_of(inv["clean_pdf"]) and inv["date"].weekday() < 5, inv["no"]
    assert "98,800" in text_of(new_lease, 2) and "4." in text_of(new_lease, 2) and len(fitz.open(stream=new_lease)) == 3
    sl = contracts["docs/contracts/S02_side_letter_2024.pdf"]
    assert "rate of 1.50% per annum" in text_of(sl).replace("\n", " ")
    assert "2.00% p.a." in text_of(INBOX[2]["clean_pdf"])

    # ---- staged history
    W = lambda s: datetime.fromisoformat(s)  # noqa: E731
    hashes = {}
    write(".gitignore", "ground_truth/\n~$*\n")
    write("COMPANY.md", company_md(False))
    write_coa()
    write_master(W("2026-01-09T10:12:00"))
    for k, v in contracts.items():
        write(k, v)
    write_fx("2026-01", W("2026-01-09T10:12:00"))
    write_commitments(D(2026, 1, 9), W("2026-01-09T10:12:00"))
    hashes["setup"] = commit("Set up family office workspace: COMPANY.md, chart of accounts, counterparty master, "
                             "fund commitments, contracts", "2026-01-09T10:12:00+08:00", "grace")
    write_bva(0, W("2026-01-23T16:40:00"))
    hashes["budget"] = commit("Budget 2026 approved (V. Cheung, 20 Jan)", "2026-01-23T16:40:00+08:00", "grace")
    write_gl(D(2026, 6, 30))
    write_payments(D(2026, 6, 30))
    write_bva(6, W("2026-07-10T11:05:00"))
    write_fx("2026-07", W("2026-07-10T11:05:00"))
    write_commitments(D(2026, 6, 30), W("2026-07-10T11:05:00"))
    hashes["h1"] = commit("H1 close: GL export Jan-Jun, payments history, budget vs actual and commitments to June",
                          "2026-07-10T11:05:00+08:00", "grace")

    def enter(until, when_s, msg, who, gl_until=None, bva_m=None, fx_m=None):
        when = W(when_s[:19])
        write_register(until, when.date(), when)
        write_commitments(until, when)
        for inv in HIST:
            if inv["date"] <= until:
                rel = f"{inv['folder']}/{inv['file']}"
                write(rel, pdfs[rel])
        if gl_until:
            write_gl(gl_until)
            write_payments(gl_until)
            write_bva(bva_m, when)
            write_fx(fx_m, when)
        return commit(msg, when_s, who)

    hashes["jul"] = enter(D(2026, 7, 31), "2026-08-05T15:20:00+08:00",
                          "Enter July documents (Pearl River call 17, Harbourview Q3 fee); GL export to 31 Jul", "jason",
                          D(2026, 7, 31), 7, "2026-08")
    eml("docs/emails/2026-08-14_raymond_forecast_assumptions.eml", f"Raymond Ho <{PEOPLE['raymond'][1]}>",
        f"Grace Lam <{PEOPLE['grace'][1]}>", "2027-28 forecast assumptions", datetime(2026, 8, 14, 9, 47),
        "Grace,\n\nFor the 2027-28 forecast please keep the Repulse Bay rent (Flat 12A, Halcyon Re) flat at HK$95,000/month\n"
        "until the lease is renewed. Halcyon have said they want to stay but nothing is agreed yet. Victoria is fine with this.\n"
        "Portfolio income +3% in 2028, staff flat.\n\nRaymond\n", cc=f"Victoria Cheung <{PEOPLE['victoria'][1]}>")
    write("COMPANY.md", company_md(True))
    write_forecast(W("2026-08-14T17:30:00"))
    hashes["fc"] = commit("Forecast 2027-28 v1: rental income flat pending lease renewal", "2026-08-14T17:30:00+08:00", "grace")
    hashes["aug"] = enter(D(2026, 8, 31), "2026-09-02T14:10:00+08:00",
                          "Enter August documents (Harbourview call 3/2026); GL export to 31 Aug", "jason",
                          D(2026, 8, 31), 8, "2026-09")
    hashes["sep1"] = enter(D(2026, 9, 3), "2026-09-03T11:45:00+08:00",
                           "Enter early September invoices (incl. Peak Estates INV-PM-2291)", "jason")
    hashes["sep2"] = enter(D(2026, 9, 18), "2026-09-18T16:00:00+08:00", "Enter mid-September invoices", "jason")
    write_register(CLOSE, CLOSE, W("2026-10-02T10:30:00"))
    write_commitments(CLOSE, W("2026-10-02T10:30:00"))
    write_gl(CLOSE)
    write_payments(CLOSE)
    write_bva(9, W("2026-10-02T10:30:00"))
    write_fx("2026-10", W("2026-10-02T10:30:00"))
    write("docs/bank/statement_2026-09.pdf", stmt_pdf)
    hashes["close"] = commit("September close (in progress): bank statement, payments run, GL export to 30 Sep",
                             "2026-10-02T10:30:00+08:00", "grace")

    # sidecar sources (backfilled for key historical cells)
    regs = reg_rows(CLOSE, CLOSE)
    rowno = {r[3]: i + 2 for i, r in enumerate(regs)}
    s02_q3 = by_no["HCP3-MF-2026Q3"]
    old_lease_p = find_page(contracts["docs/contracts/S04_lease_2024_2026.pdf"], "HK$95,000")
    s01_c = contracts["docs/contracts/S01_subscription_agreement_2024.pdf"]
    sl_p = find_page(sl, "3.1")
    sources = [
        dict(file="sheets/forecast_2027_2028.xlsx", sheet="Forecast", cell=f"C{RENT_ROW}:Z{RENT_ROW}", value=95000,
             sources=[dict(doc="COMPANY.md", page=None, quote=RENT_ASSUMPTION),
                      dict(doc="docs/emails/2026-08-14_raymond_forecast_assumptions.eml", page=None,
                           quote="please keep the Repulse Bay rent (Flat 12A, Halcyon Re) flat at HK$95,000/month"),
                      dict(doc="docs/contracts/S04_lease_2024_2026.pdf", page=old_lease_p,
                           quote="The monthly rent shall be HK$95,000 (Hong Kong Dollars Ninety-Five Thousand) inclusive of management fees and exclusive of rates, fixed for the whole Term.")],
             reason="Repulse Bay rent held flat at the current lease rate until the lease expiring 31 Dec 2026 is renewed.",
             commit=hashes["fc"]),
        dict(file="sheets/invoice_register.xlsx", sheet="Register", cell=f"F{rowno['INV-PM-2291']}", value=INV2291["total"],
             sources=[dict(doc=f"{INV2291['folder']}/{INV2291['file']}", page=1, quote=f"TOTAL HKD {money(INV2291['total'])}")],
             reason="Entered Peak Estates INV-PM-2291 (September management fees, inspections and cleaning).", commit=hashes["sep1"]),
        dict(file="sheets/invoice_register.xlsx", sheet="Register", cell=f"F{rowno[s02_q3['no']]}", value=s02_q3["total"],
             sources=[dict(doc=f"{s02_q3['folder']}/{s02_q3['file']}", page=1, quote=f"Fee rate: {FEE_RATE:.2f}% p.a. on commitment of USD 10,000,000"),
                      dict(doc="docs/contracts/S02_side_letter_2024.pdf", page=sl_p, clause="3.1",
                           quote=f"calculated at the rate of {FEE_RATE:.2f}% per annum of the Commitment")],
             reason="Entered Harbourview Q3 2026 management fee; rate matches side letter clause 3.1.", commit=hashes["jul"]),
        dict(file="sheets/supplier_master.xlsx", sheet="Suppliers", cell="F2", value=SUP["S01"]["bank_acct"],
             sources=[dict(doc="docs/contracts/S01_subscription_agreement_2024.pdf", page=find_page(s01_c, "2049"),
                           quote=f"The Fund's designated account ... account no. {SUP['S01']['bank_acct']}")],
             reason="Fund bank account per signed subscription agreement clause 5.2.", commit=hashes["setup"]),
        dict(file="sheets/budget_vs_actual_2026.xlsx", sheet="Budget monthly", cell="C" + str(PL_CODES.index("260") + 2),
             value=157000,
             sources=[dict(doc="docs/contracts/S04_lease_2024_2026.pdf", page=old_lease_p, quote="The monthly rent shall be HK$95,000")],
             reason="Budget 2026 rental income = Flat 12A lease HK$95,000 + Flat 21B HK$62,000.", commit=hashes["budget"]),
    ]
    write(".trace/sources.json", json.dumps(sources, indent=2, ensure_ascii=False) + "\n")
    hashes["src"] = commit("Backfill Trace source index for key cells", "2026-10-02T17:40:00+08:00", "grace")

    # ---- uncommitted demo material
    for inv in INBOX:
        rel = f"{inv['folder']}/{inv['file']}"
        write(rel, pdfs[rel])
    write("docs/contracts/S04_lease_2027_signed.pdf", new_lease)
    eml("docs/emails/2026-09-29_peakestates_lease_signed.eml", "Sandy Kwok <leasing@peakestates.com.hk>",
        f"Grace Lam <{PEOPLE['grace'][1]}>", "Countersigned tenancy agreement - Flat 12A Seaview Court, 2027-2029",
        datetime(2026, 9, 29, 16, 12),
        "Dear Grace,\n\nHalcyon Re have countersigned the renewal for Flat 12A, Seaview Court, for 1 Jan 2027 - 31 Dec 2029.\n"
        "Rent and the annual review are set out in clause 4. The signed copy is attached; the original follows by courier.\n"
        "Our leasing commission invoice will follow separately.\n\nBest regards,\nSandy Kwok\nLeasing, Peak Estates Property Management Ltd\n",
        cc=f"Raymond Ho <{PEOPLE['raymond'][1]}>", extra={"X-Attachment": "S04_lease_2027_signed.pdf (saved to docs/contracts/)"})
    eml("docs/emails/2026-10-02_prg_bank_change.eml", f"Pearl River Fund Accounts <accounts@{FAKE_DOMAIN}>",
        f"Lantau Peak Accounts <{PEOPLE['jason'][1]}>", "Updated wire instructions / 银行账户变更通知 - URGENT",
        datetime(2026, 10, 2, 8, 55),
        "Dear Limited Partner,\n\nPlease note the Fund's bank account has changed due to an audit of our custodian arrangements.\n"
        "Please pay drawdown notice PRG2-DN-018 and all future drawdowns to the new account:\n"
        f"{NEW_BANK[0]}, account no. {NEW_BANK[1]}.\n"
        "请将本次缴款汇入以上新账户，旧账户已停止使用。Kindly process before the due date to avoid default interest.\n\n"
        "Lily Zhang 张丽\nFund Accounting, Pearl River Capital Management\n",
        extra={"X-Attachment": "scan_1002.pdf (saved to docs/invoices/inbox/)"})

    # ---- ground truth
    s01_in, s02_call, s02_fee, dup_in, s03_in, s05_in = INBOX
    last6 = s01[-6:]

    def reg_row(inv, status, controls=()):
        f = fx(inv["currency"], inv["date"])
        return dict(date=inv["date"].isoformat(), supplier_id=inv["sid"], supplier=SUP[inv["sid"]]["name_en"], invoice_no=inv["no"],
                    currency=inv["currency"], amount=inv["total"], fx_rate=f, amount_hkd=r2(inv["total"] * f),
                    account_code=inv["account"], due_date=inv["due"].isoformat(), status=status,
                    source_file=f"{inv['folder']}/{inv['file']}", kind=inv["kind"],
                    needs_principal_approval=r2(inv["total"] * f) > APPROVAL_LIMIT, controls_fired=list(controls))

    prop = {c: dict(budget=sum(budget(c, m) for m in range(1, 10)), actual=r2(sum(actual(c, m) for m in range(1, 10))))
            for c in ("471", "473", "475")}
    for v in prop.values():
        v["variance"] = r2(v["actual"] - v["budget"])
    rm_small = r2(sum(actual("473", m) for m in range(1, 10)) - 38500 - 9800)
    pm_extra = r2(prop["471"]["actual"] - prop["471"]["budget"])
    drivers = [
        dict(driver="Emergency water-pipe repair after typhoon, Flat 12A Seaview Court (Swiftfix Building Services), 14 Jul 2026",
             account="473", amount=38500.0),
        dict(driver="Air-con replacement, Flat 21B Pinecrest Tower (Arctic Breeze Engineering), 16 Sep 2026", account="473", amount=9800.0),
        dict(driver="Building management fees revised HK$11,600 -> HK$12,800/month from Jun 2026", account="475", amount=4800.0),
        dict(driver="Routine small repairs (Island Wide Handyman) vs HK$3,000/month budget", account="473", amount=r2(rm_small - 27000)),
        dict(driver="Peak Estates extra inspections and cleaning vs HK$18,000/month budget", account="471", amount=pm_extra),
    ]
    # September cash flow, from the GL: every journal that touches the bank account (090) is one cash movement, grouped by
    # what it settles. Bank charges inside a payment journal (404) are split out.
    cf_open = balance("090", D(2026, 8, 31))
    cf_lines = {}

    def cf_add(cat, amount, jno):
        line = cf_lines.setdefault(cat, dict(amount=0.0, journals=[]))
        line["amount"] = r2(line["amount"] + amount)
        line["journals"].append(jno)

    for j in JOURNALS:
        if j["date"].month != 9 or j["date"].year != 2026:
            continue
        cash = r2(sum(dr - cr for a, dr, cr in j["lines"] if a == "090"))
        if not cash:
            continue
        fees = r2(sum(dr for a, dr, cr in j["lines"] if a == "404"))
        if fees and j["cat"] != "Bank charges":
            cf_add("Bank charges", -fees, j["no"])
            cash = r2(cash + fees)
        if cash:
            cf_add(j["cat"] or "Other", cash, j["no"])
    cf_close = r2(cf_open + sum(v["amount"] for v in cf_lines.values()))
    assert abs(cf_open - opening) < 0.005 and abs(cf_close - stmt_close) < 0.005, "cash flow does not reconcile to bank statement"
    cash_flow = dict(
        period="Sep 2026", currency="HKD", bank_account="090", opening_cash=cf_open,
        receipts={k: v for k, v in cf_lines.items() if v["amount"] > 0},
        payments={k: v for k, v in sorted(cf_lines.items(), key=lambda kv: kv[1]["amount"]) if v["amount"] < 0},
    )
    cash_flow.update(total_receipts=r2(sum(v["amount"] for v in cash_flow["receipts"].values())),
                     total_payments=r2(sum(v["amount"] for v in cash_flow["payments"].values())))
    cash_flow.update(net_cash_flow=r2(cash_flow["total_receipts"] + cash_flow["total_payments"]), closing_cash=cf_close,
                     reconciles_to=dict(doc="docs/bank/statement_2026-09.pdf", opening=opening, closing=stmt_close),
                     output_file="sheets/cash_flow_2026-09.xlsx",
                     notes=["Built from GL account 090 lines; journal numbers are the 'Journal No.' column of ledger/gl_export_2026.csv",
                            "TT charges (404) inside payment journals are shown under Bank charges, not the counterparty line",
                            "Staff & MPF = net salaries paid plus MPF remitted (both employer and employee parts)"])
    commit_rows = {r[1]: r for r in commitments(CLOSE) if r[1]}
    hcp = commit_rows["S02"]
    expected = dict(
        as_of=TODAY.isoformat(),
        planted_problems=[
            dict(id="P1", control="PRICE-001", invoice_file=f"docs/invoices/inbox/{s02_fee['file']}", invoice_no=s02_fee["no"], supplier_id="S02",
                 evidence=dict(kind="fee_rate", invoice_fee_rate_pct=FEE_RATE_WRONG, contract_fee_rate_pct=FEE_RATE, currency="USD",
                               pct_over=round((FEE_RATE_WRONG - FEE_RATE) / FEE_RATE * 100, 2), tolerance_pct=1.0,
                               invoice_amount=s02_fee["total"], expected_amount=r2(SUP["S02"]["commitment"] * FEE_RATE / 100 / 4),
                               overcharge_usd=r2(s02_fee["total"] - SUP["S02"]["commitment"] * FEE_RATE / 100 / 4),
                               contract_ref=dict(doc="docs/contracts/S02_side_letter_2024.pdf", page=sl_p, clause="3.1"))),
            dict(id="P2", control="BANK-001", invoice_file=f"docs/invoices/inbox/{s01_in['file']}", invoice_no=s01_in["no"], supplier_id="S01",
                 evidence=dict(invoice_bank_account=NEW_BANK[1], invoice_last4="7731", master_bank_account=SUP["S01"]["bank_acct"],
                               master_last4="2049", last_6_payments=[dict(date=p["payment_date"].isoformat(), invoice_no=p["invoice_no"],
                                                                          amount=p["amount"], account=p["bank_account_paid"]) for p in last6],
                               email="docs/emails/2026-10-02_prg_bank_change.eml",
                               contract_clause=dict(doc="docs/contracts/S01_subscription_agreement_2024.pdf", page=find_page(s01_c, "5.3"), clause="5.3"))),
            dict(id="P3", control="DOMAIN-001", invoice_file=f"docs/invoices/inbox/{s01_in['file']}", invoice_no=s01_in["no"], supplier_id="S01",
                 evidence=dict(sender_domain=FAKE_DOMAIN, known_domain=SUP["S01"]["domain"],
                               email="docs/emails/2026-10-02_prg_bank_change.eml", invoice_contact_email=f"accounts@{FAKE_DOMAIN}")),
            dict(id="P4", control="DUP-001", invoice_file=f"docs/invoices/inbox/{dup_in['file']}", invoice_no=dup_in["no"], supplier_id="S03",
                 evidence=dict(matches_invoice_no="INV-PM-2291", original_date=INV2291["date"].isoformat(), entered="2026-09-03",
                               amount=INV2291["total"], currency="HKD", register_row=rowno["INV-PM-2291"],
                               original_status="Approved (due 2026-10-01, unpaid)")),
        ],
        clean_invoices=[f"docs/invoices/inbox/{i['file']}" for i in (s02_call, s03_in, s05_in)],
        register_seed_rows=len(regs),
        tasks=[
            dict(id=1, prompt="Halcyon Re signed the Repulse Bay renewal. Update the forecast.",
                 expected=dict(file="sheets/forecast_2027_2028.xlsx", named_range=RENT_RANGE, sheet="Forecast",
                               cells={f"{get_column_letter(3 + j)}{RENT_ROW}": (98800.0 if j < 12 else 101764.0) for j in range(24)},
                               changed_cells=24, source=dict(doc="docs/contracts/S04_lease_2027_signed.pdf", page=3, clause="4"),
                               also=["COMPANY.md rental income assumption should be updated (lease renewed, +3%/yr)",
                                     "supplier_master S04 contract_file -> S04_lease_2027_signed.pdf"])),
            dict(id=2, prompt="Enter this week's capital calls, fee notices and invoices.",
                 expected=dict(rows=[reg_row(s02_call, "Entered"), reg_row(s03_in, "Entered"), reg_row(s05_in, "Entered"),
                                     reg_row(s01_in, "HELD", ["BANK-001", "DOMAIN-001"]), reg_row(s02_fee, "HELD", ["PRICE-001"]),
                                     reg_row(dup_in, "HELD", ["DUP-001"])])),
            dict(id=3, prompt="Why is 2027 rental income for Flat 12A 98,800?",
                 expected=dict(facts=["Renewal tenancy agreement LPFO/L/2027/12A with Halcyon Re Asia Ltd signed 29 Sep 2026",
                                      "Clause 4.1 (page 3): rent HK$98,800/month from 1 Jan 2027",
                                      "Clause 4.2: +3% every 1 January (HK$101,764 in 2028)",
                                      "Previously HK$95,000 flat: old lease (expires 31 Dec 2026) and COMPANY.md assumption by R. Ho, Aug 2026",
                                      "Increase vs prior forecast: +HK$3,800/month (+4.0%)"],
                               sources=[dict(doc="docs/contracts/S04_lease_2027_signed.pdf", page=3, clause="4"),
                                        dict(doc="docs/emails/2026-09-29_peakestates_lease_signed.eml"),
                                        dict(doc="COMPANY.md"), dict(doc="docs/emails/2026-08-14_raymond_forecast_assumptions.eml"),
                                        dict(commit=hashes["fc"], note="forecast v1 rental income flat")])),
            dict(id=4, prompt="Which contracts need action in the next 30 days?",
                 expected=dict(items=[dict(contract="docs/contracts/S05_art_insurance_policy_2025.pdf", supplier_id="S05",
                                           renewal_date="2026-11-16", notice_days=30, notice_deadline="2026-10-17",
                                           commitment_if_missed="12 months x USD 2,450 = USD 29,400")],
                               not_in_window=["S04 old lease expires 2026-12-31 (already replaced by the renewal)",
                                              "S03 property management agreement expires 2026-12-31"])),
            dict(id=5, prompt="Why are property costs over budget YTD?",
                 expected=dict(period="Jan-Sep 2026", accounts=prop,
                               property_budget=sum(v["budget"] for v in prop.values()),
                               property_actual=r2(sum(v["actual"] for v in prop.values())),
                               property_variance=r2(sum(v["variance"] for v in prop.values())), drivers=drivers)),
            dict(id=6, prompt="Prepare September's cash flow report.", expected=cash_flow),
            dict(id=7, prompt="How much do we still owe Harbourview Capital Partners III?",
                 expected=dict(file="sheets/commitments.xlsx", fund=hcp[0], currency="USD", commitment=hcp[3],
                               called_to_date=hcp[4], unfunded=r2(hcp[3] - hcp[4]), as_of=CLOSE.isoformat(),
                               pending=dict(doc=f"docs/invoices/inbox/{s02_call['file']}", amount=s02_call["total"],
                                            unfunded_after=r2(hcp[3] - hcp[4] - s02_call["total"])))),
        ],
    )
    write("ground_truth/expected.json", json.dumps(expected, indent=2, ensure_ascii=False) + "\n")

    # ---- re-read checks on written files
    wb = load_workbook(WS / "sheets/budget_vs_actual_2026.xlsx")
    gl_sum = {}
    with open(WS / "ledger/gl_export_2026.csv") as f:
        for r in csv.DictReader(f):
            gl_sum[r["Account Code"]] = gl_sum.get(r["Account Code"], 0) + float(r["Net"])
    for row in wb["YTD"].iter_rows(min_row=5, max_row=4 + len(PL_CODES), values_only=True):
        g = gl_sum.get(row[0], 0) * (-1 if ACC[row[0]][2] == "Revenue" else 1)
        assert abs(g - row[4]) < 0.01, (row, g)
    wbr = load_workbook(WS / "sheets/invoice_register.xlsx")["Register"]
    assert abs(sum(r[5] for r in wbr.iter_rows(min_row=2, values_only=True)) - sum(i["total"] for i in HIST)) < 0.01

    # ---- summary
    print("workspace/ generated")
    for d in sorted({p.parent for p in WS.rglob("*") if p.is_file() and GIT_DIR not in p.parents}):
        files = sorted(f.name for f in d.iterdir() if f.is_file())
        print(f"  {d.relative_to(WS)}/  ({len(files)} files)")
    print(f"GL rows {len(GL)} ({len(JOURNALS)} journals); register rows {len(HIST)}; payments {len(PAYMENTS)}; "
          f"Sep bank lines {len(stmt_lines)}; opening {money(opening)} closing {money(stmt_close)}")
    for code in ("471", "473", "475", "481", "455"):
        b = sum(budget(code, m) for m in range(1, 10))
        a = sum(actual(code, m) for m in range(1, 10))
        print(f"  {code} {ACC[code][1]:<38} budget {b:>10,.0f} actual {a:>12,.2f} var {a - b:>+10,.0f}")
    print(git("log", "--format=%h %ad %an  %s", "--date=short"))
    print(git("status", "--short"))


if __name__ == "__main__":
    main()
