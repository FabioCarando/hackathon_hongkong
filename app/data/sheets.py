"""Read and write the team's xlsx sheets with openpyxl, keeping formulas and formatting."""

import re
from copy import copy
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter, range_boundaries

from app.data import workspace


def _cell_value(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    return v


def read_table(file: str, sheet: str) -> list[dict[str, Any]]:
    """Rows as dicts (header = row 1) with the sheet row number in `_row`."""
    ws = load_workbook(workspace.path(file))[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    out = []
    for i, r in enumerate(rows[1:], start=2):
        if all(v is None for v in r):
            continue
        out.append({"_row": i, **{h: _cell_value(v) for h, v in zip(header, r)}})
    return out


def register() -> list[dict[str, Any]]:
    """invoice_register rows with amount_hkd computed (the sheet stores a formula)."""
    rows = read_table("sheets/invoice_register.xlsx", "Register")
    for r in rows:
        if isinstance(r.get("amount_hkd"), str):
            r["amount_hkd"] = round(float(r["amount"]) * float(r["fx_rate"]), 2)
    return rows


def suppliers() -> list[dict[str, Any]]:
    return read_table("sheets/supplier_master.xlsx", "Suppliers")


def fx_rate(currency: str, on: date) -> float:
    if currency == "HKD":
        return 1.0
    month = on.strftime("%Y-%m")
    rows = read_table("sheets/fx_rates.xlsx", "Monthly")
    col = {"USD": "USD_HKD", "RMB": "RMB_HKD"}[currency]
    match = [r for r in rows if r["month"] == month] or rows[-1:]
    return float(match[0][col])


def read_range(file: str, sheet: str, ref: str) -> dict[str, Any]:
    ws = load_workbook(workspace.path(file))[sheet]
    c1, r1, c2, r2 = range_boundaries(ref if ":" in ref else f"{ref}:{ref}")
    return {
        f"{get_column_letter(c)}{r}": _cell_value(ws.cell(r, c).value)
        for r in range(r1, r2 + 1)
        for c in range(c1, c2 + 1)
    }


def cells_in(ref: str) -> list[str]:
    c1, r1, c2, r2 = range_boundaries(ref if ":" in ref else f"{ref}:{ref}")
    return [f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def write_cells(file: str, sheet: str, values: dict[str, Any], comment: str | None = None) -> dict:
    """Write {cell: value}; returns the old values."""
    p = workspace.path(file)
    wb = load_workbook(p)
    ws = wb[sheet]
    old = {}
    for cell, v in values.items():
        old[cell] = _cell_value(ws[cell].value)
        ws[cell].value = v
        if comment:
            ws[cell].comment = Comment(comment, "Trace")
    wb.save(p)
    return old


def append_row(file: str, sheet: str, values: dict[str, Any], comment: str | None = None) -> int:
    """Append after the last used row, copying the previous row's style and formula pattern."""
    p = workspace.path(file)
    wb = load_workbook(p)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    last = max(r for r in range(1, ws.max_row + 1) if any(c.value is not None for c in ws[r]))
    row = last + 1
    for col, name in enumerate(header, start=1):
        prev = ws.cell(last, col)
        cell = ws.cell(row, col)
        if prev.has_style:
            cell._style = copy(prev._style)
        v = values.get(name)
        if isinstance(prev.value, str) and prev.value.startswith("="):
            v = prev.value.replace(str(last), str(row))  # e.g. =ROUND(F18*G18,2)
        elif isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            v = datetime.fromisoformat(v)  # dates come back from JSON as ISO strings
        elif isinstance(v, date) and not isinstance(v, datetime):
            v = datetime(v.year, v.month, v.day)
        cell.value = v
    if comment:
        ws.cell(row, 1).comment = Comment(comment, "Trace")
    wb.save(p)
    return row


def header(file: str, sheet: str) -> list[str]:
    ws = load_workbook(workspace.path(file))[sheet]
    return [c.value for c in ws[1]]


def named_range(file: str, name: str) -> tuple[str, str] | None:
    """(sheet, 'C5:Z5') for a workbook-level defined name, or None."""
    wb = load_workbook(workspace.path(file))
    dn = wb.defined_names.get(name)
    if dn is None:
        return None
    sheet, ref = next(iter(dn.destinations))
    return sheet, ref.replace("$", "")


def find_row(file: str, sheet: str, contains: str, column: str = "A") -> int | None:
    """First row whose cell in `column` contains the text (case-insensitive)."""
    ws = load_workbook(workspace.path(file))[sheet]
    for row in range(1, ws.max_row + 1):
        v = ws[f"{column}{row}"].value
        if isinstance(v, str) and contains.lower() in v.lower():
            return row
    return None
